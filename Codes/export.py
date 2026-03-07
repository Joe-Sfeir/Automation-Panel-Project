#!/usr/bin/env python3
"""
export.py  —  Export a filtered offer from a matched offer Excel.

Reads the "Pricing Detail" sheet (written by test_match.py / run_match()),
lets you choose which columns to include, and writes a clean formatted Excel.

─────────────────────────────────────────────────────────────
AVAILABLE COLUMNS  (pass any subset to export_offer)
─────────────────────────────────────────────────────────────
  "desc"        Description (client-facing name)
  "brand"       Brand / manufacturer
  "qty"         Quantity
  "ref"         Manufacturer reference number
  "unit_price"  Unit price (EUR)
  "total"       Line total (qty × unit price)
  "ppl_desc"    Exact PPL product description (internal review)
  "tier"        Match tier: T1/T2/T3/NF  (internal review)

─────────────────────────────────────────────────────────────
PYSIDE6 USAGE
─────────────────────────────────────────────────────────────
    from export import AVAILABLE_COLUMNS, export_offer

    # Show user a checkbox list of AVAILABLE_COLUMNS.values()
    # Collect which keys they ticked, e.g.:
    selected = ["desc", "brand", "qty", "ref"]

    path = export_offer(
        source_excel = "DB-B1_offer.xlsx",
        columns      = selected,
        out_path     = "DB-B1_export.xlsx",
    )
    print(f"Exported to {path}")

─────────────────────────────────────────────────────────────
CLI USAGE
─────────────────────────────────────────────────────────────
    python export.py --source DB-B1_offer.xlsx \\
                     --columns desc brand qty ref unit_price \\
                     --out DB-B1_client.xlsx
"""

import argparse
import sys
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


# ─────────────────────────────────────────────────────────────────────────────
#  COLUMN REGISTRY
#  Key   = internal ID your UI uses
#  Value = human label shown in the checkbox list
# ─────────────────────────────────────────────────────────────────────────────

AVAILABLE_COLUMNS: dict[str, str] = {
    "desc":       "Description",
    "brand":      "Brand",
    "qty":        "Qty",
    "ref":        "Ref.",
    "unit_price": "Unit Price (EUR)",
    "total":      "Total (EUR)",
    "ppl_desc":   "PPL Description",
    "tier":       "Match Tier",
}

# Maps column key → column header in the Pricing Detail sheet
_PD_HEADER_MAP: dict[str, str] = {
    "desc":       "Description",
    "brand":      "Brand",
    "qty":        "Qty",
    "ref":        "Ref.",
    "unit_price": "Unit Price (EUR)",
    "total":      "Total (EUR)",
    "ppl_desc":   "PPL Description",
    "tier":       "Match Tier",
}

# Column widths for each key
_COL_WIDTHS: dict[str, float] = {
    "desc":       44.0,
    "brand":      15.0,
    "qty":         6.0,
    "ref":        18.0,
    "unit_price": 16.0,
    "total":      16.0,
    "ppl_desc":   52.0,
    "tier":        9.0,
}

# Tier colours (same as Pricing Detail)
_TIER_CLR = {"T1": "C6EFCE", "T2": "FFEB9C", "T3": "BDD7EE", "NF": "D9D9D9"}

_HDR_FILL = PatternFill("solid", fgColor="1F3864")
_HDR_FONT = Font(name="Arial", bold=True, size=10, color="FFFFFF")
_DAT_FONT = Font(name="Arial", size=10)
_THIN     = Side(style="thin", color="BBBBBB")
_BORDER   = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


# ─────────────────────────────────────────────────────────────────────────────
#  READER  —  parse Pricing Detail into list of row dicts
# ─────────────────────────────────────────────────────────────────────────────

def _read_pricing_detail(source_excel: str | Path) -> tuple[list[dict], list[str]]:
    """
    Returns (rows, panel_order).
    rows        — list of dicts with keys matching AVAILABLE_COLUMNS + "panel"
    panel_order — panel names in original order
    """
    wb = openpyxl.load_workbook(str(source_excel), data_only=True)

    if "Pricing Detail" not in wb.sheetnames:
        raise ValueError(f'"Pricing Detail" sheet not found in {source_excel}.\n'
                         f'Run test_match.py first to generate it.')

    ws = wb["Pricing Detail"]

    # Build header → col index map from row 1
    hdr_to_col: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(1, col).value
        if val:
            hdr_to_col[str(val).strip()] = col

    # Reverse map: internal key → column index in this sheet
    key_to_col: dict[str, int] = {}
    for key, hdr in _PD_HEADER_MAP.items():
        if hdr in hdr_to_col:
            key_to_col[key] = hdr_to_col[hdr]

    panel_col = hdr_to_col.get("Panel")
    if not panel_col:
        raise ValueError('"Panel" column not found in Pricing Detail sheet.')

    rows        = []
    panel_order = []
    seen_panels = set()

    for r in range(2, ws.max_row + 1):
        panel_val = ws.cell(r, panel_col).value
        if not panel_val:
            continue  # legend / blank row
        panel = str(panel_val).strip()
        if panel not in seen_panels:
            panel_order.append(panel)
            seen_panels.add(panel)

        row: dict = {"panel": panel}
        for key, col in key_to_col.items():
            row[key] = ws.cell(r, col).value
        rows.append(row)

    wb.close()
    return rows, panel_order


# ─────────────────────────────────────────────────────────────────────────────
#  WRITER  —  produce clean export Excel
# ─────────────────────────────────────────────────────────────────────────────

def _write_export(rows: list[dict], panel_order: list[str],
                  columns: list[str], out_path: Path):
    """Write a clean grouped Excel with only the requested columns."""
    wb  = Workbook()
    ws  = wb.active
    ws.title = "Export"

    # Validate columns
    valid = [c for c in columns if c in AVAILABLE_COLUMNS]
    if not valid:
        raise ValueError(f"No valid columns selected. Choose from: {list(AVAILABLE_COLUMNS)}")

    # ── Column headers ────────────────────────────────────────────────────────
    # First col = Panel (always, for grouping)
    all_cols  = ["panel"] + valid
    disp_hdrs = ["Panel"] + [AVAILABLE_COLUMNS[c] for c in valid]

    for ci, (col_key, hdr) in enumerate(zip(all_cols, disp_hdrs), 1):
        cell            = ws.cell(1, ci, hdr)
        cell.font       = _HDR_FONT
        cell.fill       = _HDR_FILL
        cell.alignment  = Alignment(horizontal="center", vertical="center")
        cell.border     = _BORDER
        w = 18.0 if col_key == "panel" else _COL_WIDTHS.get(col_key, 14.0)
        ws.column_dimensions[get_column_letter(ci)].width = w

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    # ── Data rows, grouped by panel ──────────────────────────────────────────
    by_panel: dict[str, list] = {p: [] for p in panel_order}
    for row in rows:
        by_panel[row["panel"]].append(row)

    dr = 2
    for panel in panel_order:
        panel_rows = by_panel[panel]
        if not panel_rows:
            continue

        # Panel sub-header row
        ws.merge_cells(start_row=dr, start_column=1,
                       end_row=dr, end_column=len(all_cols))
        ph_cell            = ws.cell(dr, 1, panel)
        ph_cell.font       = Font(name="Arial", bold=True, size=11, color="FFFFFF")
        ph_cell.fill       = PatternFill("solid", fgColor="2E4057")
        ph_cell.alignment  = Alignment(horizontal="left", vertical="center",
                                        indent=1)
        ws.row_dimensions[dr].height = 18
        dr += 1

        for item in panel_rows:
            tier      = str(item.get("tier") or "")
            fill_clr  = _TIER_CLR.get(tier, "FFFFFF")
            fill_c    = PatternFill("solid", fgColor=fill_clr)

            for ci, col_key in enumerate(all_cols, 1):
                val  = item.get(col_key, "")
                cell = ws.cell(dr, ci, val)
                cell.font   = _DAT_FONT
                cell.fill   = fill_c
                cell.border = _BORDER

                # Alignment and number format
                if col_key in ("unit_price", "total"):
                    cell.number_format = '#,##0.00'
                    cell.alignment     = Alignment(horizontal="right",
                                                   vertical="center")
                elif col_key in ("qty", "ref", "brand", "tier", "panel"):
                    cell.alignment = Alignment(horizontal="center",
                                              vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left",
                                              vertical="center",
                                              wrap_text=(col_key == "ppl_desc"))

            ws.row_dimensions[dr].height = 15
            dr += 1

        # Totals row (only if unit_price or total is selected)
        if "total" in valid or "unit_price" in valid:
            panel_total = sum(
                (r.get("total") or 0) for r in panel_rows
            )
            ws.cell(dr, 1, "TOTAL").font      = Font(name="Arial", bold=True, size=10)
            ws.cell(dr, 1).alignment          = Alignment(horizontal="right")
            ws.cell(dr, 1).fill               = PatternFill("solid", fgColor="D9E1F2")

            if "total" in valid:
                tc = all_cols.index("total") + 1
                cell                = ws.cell(dr, tc, panel_total)
                cell.font           = Font(name="Arial", bold=True, size=10)
                cell.number_format  = '#,##0.00'
                cell.alignment      = Alignment(horizontal="right", vertical="center")
                cell.fill           = PatternFill("solid", fgColor="D9E1F2")

            dr += 1

        # Blank separator between panels
        dr += 1

    # ── Legend (only if tier column shown) ───────────────────────────────────
    if "tier" in valid:
        dr += 1
        for text, color in [("T1 = Keyword match", "C6EFCE"),
                             ("T2 = AI matched",    "FFEB9C"),
                             ("T3 = Substituted",   "BDD7EE"),
                             ("NF = Not found",     "D9D9D9")]:
            cell      = ws.cell(dr, 1, text)
            cell.fill = PatternFill("solid", fgColor=color)
            cell.font = Font(name="Arial", size=9, italic=True)
            dr += 1

    wb.save(str(out_path))


# ─────────────────────────────────────────────────────────────────────────────
#  PUBLIC FUNCTION  —  called by PySide6 UI
# ─────────────────────────────────────────────────────────────────────────────

def export_offer(
    source_excel: str | Path,
    columns:      list[str],
    out_path:     str | Path,
) -> str:
    """
    Read the matched offer Excel, export only the requested columns.

    Parameters
    ----------
    source_excel : path to the offer Excel produced by run_match()
    columns      : list of column keys from AVAILABLE_COLUMNS, e.g.
                   ["desc", "brand", "qty", "ref"]
    out_path     : where to write the exported Excel

    Returns
    -------
    str — absolute path of the written file
    """
    source_excel = Path(source_excel)
    out_path     = Path(out_path)

    if not source_excel.exists():
        raise FileNotFoundError(f"Source not found: {source_excel}")

    rows, panel_order = _read_pricing_detail(source_excel)

    if not rows:
        raise ValueError("Pricing Detail sheet is empty — nothing to export.")

    _write_export(rows, panel_order, columns, out_path)
    print(f"  Exported → {out_path}  ({len(rows)} rows, columns: {columns})")
    return str(out_path.resolve())


# ─────────────────────────────────────────────────────────────────────────────
#  CLI
# ─────────────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Export a filtered offer from a matched offer Excel.",
        formatter_class=argparse.RawTextHelpFormatter,
        epilog=f"Available columns: {', '.join(AVAILABLE_COLUMNS)}"
    )
    parser.add_argument("--source",  required=True,
                        help="Offer Excel produced by test_match.py")
    parser.add_argument("--columns", nargs="+", required=True,
                        choices=list(AVAILABLE_COLUMNS),
                        metavar="COL",
                        help="Columns to include in export")
    parser.add_argument("--out",     default="export.xlsx",
                        help="Output file (default: export.xlsx)")
    args = parser.parse_args()

    try:
        result = export_offer(args.source, args.columns, args.out)
        print(f"Done → {result}")
    except (FileNotFoundError, ValueError) as e:
        print(f"ERROR: {e}"); sys.exit(1)