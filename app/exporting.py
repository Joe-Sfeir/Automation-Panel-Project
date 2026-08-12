from __future__ import annotations

import json
from datetime import date
from pathlib import Path


COLUMNS = {
    "panel": ("Panel", "panel_reference"),
    "description": ("Description", "description"),
    "manufacturer": ("Manufacturer", "manufacturer"),
    "reference": ("Catalog reference", "chosen_reference"),
    "quantity": ("Qty", "quantity"),
    "unit_price": ("Unit price", "unit_price"),
    "total_price": ("Total", "total_price"),
    "currency": ("Currency", "currency"),
    "status": ("Review status", "review_status"),
    "tier": ("Match tier", "tier"),
}


def export_pricing_json(
    pricing_json: str | Path,
    output: str | Path,
    visible_columns: list[str] | None = None,
    attention: str = "",
) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for Excel export") from exc

    payload = json.loads(Path(pricing_json).read_text(encoding="utf-8"))
    columns = [value for value in (visible_columns or list(COLUMNS)) if value in COLUMNS]
    if not columns:
        raise ValueError("At least one export column must be visible")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Offer"
    sheet.append(["TECHNICAT AUTOMATION OFFER"])
    sheet.append(["Reference", payload.get("project_reference", "")])
    sheet.append(["Attention", attention])
    sheet.append(["Date", date.today().isoformat()])
    sheet.append([])
    sheet.append([COLUMNS[column][0] for column in columns])
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(horizontal="center")
    for panel in payload.get("panels", []):
        for item in panel.get("items", []):
            item = dict(item)
            item["panel_reference"] = panel.get("panel_reference", "")
            sheet.append([item.get(COLUMNS[column][1], "") for column in columns])
    sheet.append([])
    sheet.append(["Subtotal (HT)", payload.get("grand_ht", 0.0)])
    sheet.append([f"TVA ({float(payload.get('tva_rate', 0)):.1%})", payload.get("tva", 0.0)])
    sheet.append(["Total (TTC)", payload.get("grand_ttc", 0.0)])
    for column_index, column in enumerate(columns, 1):
        width = 16
        if column in {"description", "reference"}:
            width = 38
        sheet.column_dimensions[sheet.cell(1, column_index).column_letter].width = width
    for row in sheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    destination = Path(output)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination
