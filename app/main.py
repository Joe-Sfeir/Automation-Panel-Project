from __future__ import annotations

import json
import io
import os
import re
import runpy
import shutil
import sqlite3
import subprocess
import sys
import urllib.error
import urllib.request
from contextlib import redirect_stderr, redirect_stdout
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path

from PySide6.QtCore import QObject, QRect, QSize, QThread, Qt, QUrl, Signal
from PySide6.QtGui import QColor, QDesktopServices, QIcon, QFont, QCursor, QPixmap
from PySide6.QtWidgets import (
    QApplication,
    QButtonGroup,
    QCheckBox,
    QComboBox,
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QFormLayout,
    QFrame,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QListWidget,
    QListWidgetItem,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QProgressBar,
    QProgressDialog,
    QSpinBox,
    QStackedWidget,
    QTableWidget,
    QTableWidgetItem,
    QTextEdit,
    QSizePolicy,
    QVBoxLayout,
    QWidget,
    QHeaderView,
    QAbstractItemView,
    QStyledItemDelegate,
)

from app.pipeline import ProjectPipeline
from app.persistence import CANONICAL_FILENAME
from app.review_ui import ProjectReviewDialog

SOURCE_ROOT = Path(__file__).resolve().parents[1]
IS_FROZEN = bool(getattr(sys, "frozen", False))
BUNDLE_ROOT = Path(getattr(sys, "_MEIPASS", SOURCE_ROOT))
APPDATA_ROOT = Path(os.getenv("APPDATA", str(Path.home())))
PROJECT_ROOT = (APPDATA_ROOT / "AutomationPlatform") if IS_FROZEN else SOURCE_ROOT
RUNTIME_CODES_DIR = PROJECT_ROOT / "Codes"
BUNDLED_CODES_DIR = BUNDLE_ROOT / "Codes"
CONFIG_PATH = PROJECT_ROOT / "app_config.json"
HISTORY_PATH = PROJECT_ROOT / "app_history.json"
DEFAULT_BRANDS = ["SCHNEIDER", "EATON", "LEGRAND", "ABB", "SIEMENS"]
DEFAULT_GEMINI_MODEL = "gemini-3.6-flash"


def bootstrap_runtime_workspace() -> None:
    PROJECT_ROOT.mkdir(parents=True, exist_ok=True)
    (PROJECT_ROOT / "output").mkdir(parents=True, exist_ok=True)

    if not IS_FROZEN:
        return

    RUNTIME_CODES_DIR.mkdir(parents=True, exist_ok=True)

    code_files = [
        "dwg_test3.py",
        "match_offer.py",
        "import.py",
        "import_using_all_api.py",
        "schema.sql",
    ]
    for name in code_files:
        src = BUNDLED_CODES_DIR / name
        dst = RUNTIME_CODES_DIR / name
        if src.exists():
            shutil.copy2(src, dst)

    db_src = BUNDLE_ROOT / "pricing.db"
    db_dst = PROJECT_ROOT / "pricing.db"
    if db_src.exists() and not db_dst.exists():
        shutil.copy2(db_src, db_dst)

    for dwg in BUNDLE_ROOT.glob("*.dwg"):
        dst = PROJECT_ROOT / dwg.name
        if not dst.exists():
            shutil.copy2(dwg, dst)

    prices_src = BUNDLE_ROOT / "PriceLists"
    prices_dst = PROJECT_ROOT / "PriceLists"
    if prices_src.exists() and prices_src.is_dir() and not prices_dst.exists():
        shutil.copytree(prices_src, prices_dst)

    oda_src = BUNDLE_ROOT / "ODA"
    oda_dst = PROJECT_ROOT / "ODA"
    if oda_src.exists() and oda_src.is_dir() and not oda_dst.exists():
        shutil.copytree(oda_src, oda_dst)


bootstrap_runtime_workspace()

# Revision-takeoff workbench visual system. The base declarations below remain
# for compatibility with older pages; the authored engineering overrides that
# follow are the canonical application styling.
APP_STYLESHEET = """
/* Base App Styling */
QWidget {
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Oxygen, Ubuntu, Cantarell, "Helvetica Neue", sans-serif;
    font-size: 13px;
    color: #1E293B; /* Slate 800 */
}

QMainWindow, QDialog {
    background-color: #F8FAFC; /* Slate 50 - Very light blue-gray */
}

/* Sidebar Styling - Dark Corporate Blue */
QFrame#sidebar {
    background-color: #0B1727; /* Deep Navy Blue */
    border-right: 1px solid #08101C;
}

QLabel#appLogo {
    font-size: 19px;
    font-weight: 800;
    color: #FFFFFF;
    letter-spacing: -0.5px;
}

QLabel#appLogoSubtitle {
    font-size: 11px;
    font-weight: 600;
    color: #60A5FA; /* Blue 400 */
    text-transform: uppercase;
    letter-spacing: 1px;
}

QLabel#appLogoMark {
    min-height: 84px;
    max-height: 84px;
    color: #FFFFFF;
    font-size: 24px;
    font-weight: bold;
}

QPushButton#navBtn {
    text-align: left;
    padding: 12px 16px;
    border: none;
    border-left: 4px solid transparent;
    border-radius: 6px;
    background-color: transparent;
    color: #94A3B8; /* Slate 400 */
    font-size: 14px;
    font-weight: 500;
}

QPushButton#navBtn:hover {
    background-color: #1E293B; /* Slate 800 */
    color: #FFFFFF;
}

QPushButton#navBtn:checked {
    background-color: #2563EB; /* Blue 600 */
    color: #FFFFFF;
    font-weight: 600;
    border-left: 4px solid #93C5FD; /* Light blue accent */
}

/* Typography & Headings */
QLabel#pageTitle {
    font-size: 26px;
    font-weight: 700;
    color: #0F172A; /* Slate 900 */
    letter-spacing: -0.5px;
}

QLabel#pageSubtitle {
    font-size: 14px;
    color: #64748B; /* Slate 500 */
}

QLabel#sectionTitle {
    font-size: 15px;
    font-weight: 600;
    color: #1E293B; /* Slate 800 */
    text-transform: uppercase;
    letter-spacing: 0.5px;
}

QLabel#sectionDesc {
    font-size: 13px;
    color: #64748B;
}

/* Cards & Containers */
QFrame#card {
    background-color: #FFFFFF;
    border: 1px solid #E2E8F0; /* Slate 200 */
    border-radius: 10px;
}

QFrame#actionTile {
    background-color: #FFFFFF;
    border: 1px solid #E2E8F0;
    border-radius: 12px;
}

QFrame#actionTile:hover {
    border: 1px solid #3B82F6; /* Blue 500 */
    background-color: #F0F9FF; /* Sky 50 */
}

QLabel#tileIconBubble {
    background-color: #EFF6FF; /* Blue 50 */
    color: #2563EB; /* Blue 600 */
    border-radius: 24px; /* Squircle look */
    min-width: 80px;
    max-width: 80px;
    min-height: 80px;
    max-height: 80px;
    font-size: 32px;
}

QLabel#tileIconBubble[tone="deep"] {
    background-color: #1E3A8A; /* Blue 900 */
    color: #FFFFFF;
}

QLabel#tileTitle {
    font-size: 16px;
    font-weight: 700;
    color: #0F172A;
}

QLabel#tileSubtitle {
    font-size: 12px;
    color: #64748B;
}

/* Stats */
QFrame#miniStat {
    background-color: #FFFFFF;
    border: 1px solid #E2E8F0;
    border-radius: 10px;
    border-top: 4px solid #2563EB; /* Blue accent line */
}
QLabel#statValue {
    font-size: 30px;
    font-weight: 700;
    color: #0F172A;
}
QLabel#statLabel {
    font-size: 12px;
    font-weight: 600;
    color: #64748B;
    text-transform: uppercase;
    letter-spacing: 0.5px;
}
QLabel#statMeta {
    font-size: 12px;
    color: #059669; /* Emerald 600 */
    font-weight: 500;
}

/* Buttons */
QPushButton {
    border-radius: 6px;
    padding: 8px 16px;
    font-size: 13px;
    font-weight: 600;
}

QPushButton[variant="primary"] {
    background-color: #2563EB; /* Blue 600 */
    color: white;
    border: 1px solid #1D4ED8; /* Blue 700 */
}
QPushButton[variant="primary"]:hover {
    background-color: #1D4ED8;
}
QPushButton[variant="primary"]:pressed {
    background-color: #1E40AF; /* Blue 800 */
}
QPushButton[variant="primary"]:disabled {
    background-color: #93C5FD; /* Blue 300 */
    border-color: #93C5FD;
}

QPushButton[variant="secondary"] {
    background-color: #FFFFFF;
    color: #334155; /* Slate 700 */
    border: 1px solid #CBD5E1; /* Slate 300 */
}
QPushButton[variant="secondary"]:hover {
    background-color: #F8FAFC;
    border-color: #94A3B8; /* Slate 400 */
    color: #0F172A;
}
QPushButton[variant="secondary"]:pressed {
    background-color: #F1F5F9;
}

QPushButton[variant="ghost"] {
    background-color: transparent;
    color: #2563EB;
    border: 1px solid transparent;
}
QPushButton[variant="ghost"]:hover {
    background-color: #EFF6FF;
    color: #1D4ED8;
}

/* Forms & Inputs */
QLineEdit, QSpinBox, QComboBox, QListWidget {
    background-color: #FFFFFF;
    border: 1px solid #CBD5E1;
    border-radius: 6px;
    padding: 8px 12px;
    color: #0F172A;
    selection-background-color: #BFDBFE;
}

QLineEdit:focus, QSpinBox:focus, QComboBox:focus, QListWidget:focus {
    border: 1px solid #3B82F6; /* Blue 500 */
    background-color: #FFFFFF;
}

QComboBox::drop-down {
    border: none;
    width: 24px;
}

QListWidget::item {
    padding: 8px;
    border-radius: 4px;
    margin-bottom: 2px;
}
QListWidget::item:selected {
    background-color: #EFF6FF;
    color: #1D4ED8;
    font-weight: 600;
    border-left: 3px solid #3B82F6;
}

/* Terminal / Log Output */
QTextEdit#terminal {
    background-color: #0B1727; /* Match sidebar */
    color: #F8FAFC; /* Slate 50 */
    font-family: "Cascadia Code", "Consolas", "Courier New", monospace;
    font-size: 12px;
    border: 1px solid #1E293B;
    border-radius: 8px;
    padding: 14px;
}

/* Table Widget */
QTableWidget {
    background-color: #FFFFFF;
    border: none;
    gridline-color: #F1F5F9; /* Slate 100 */
    selection-background-color: #EFF6FF;
    selection-color: #0F172A;
}
QHeaderView::section {
    background-color: #F8FAFC;
    color: #64748B;
    font-weight: 700;
    font-size: 11px;
    text-transform: uppercase;
    letter-spacing: 0.5px;
    padding: 12px 10px;
    border: none;
    border-bottom: 1px solid #E2E8F0;
    text-align: left;
}
QTableWidget::item {
    padding: 10px;
    border-bottom: 1px solid #F1F5F9;
}

/* Progress Bar */
QProgressBar {
    background-color: #E2E8F0;
    border: none;
    border-radius: 4px;
    text-align: center;
    color: transparent;
    min-height: 8px;
    max-height: 8px;
}
QProgressBar::chunk {
    background-color: #2563EB;
    border-radius: 4px;
}

/* Status Chip */
QLabel#statusChip {
    background-color: #F1F5F9;
    color: #64748B;
    padding: 4px 12px;
    border-radius: 12px;
    font-size: 11px;
    font-weight: 700;
    text-transform: uppercase;
    letter-spacing: 0.5px;
}
QLabel#statusChip[state="running"] {
    background-color: #DBEAFE; /* Blue 100 */
    color: #1D4ED8; /* Blue 700 */
}
"""

APP_STYLESHEET += """
QWidget { font-family: "Segoe UI", sans-serif; font-size: 13px; color: #20282D; }
QMainWindow, QDialog { background-color: #F2F4F3; }
QFrame#sidebar { background-color: #20292F; border-right: 1px solid #11181C; }
QLabel#appLogo { font-size: 16px; font-weight: 700; color: #F7F8F7; letter-spacing: 0.3px; }
QLabel#appLogoMark { min-height: 58px; max-height: 58px; color: #F7F8F7; font-size: 20px; font-weight: 700; }
QPushButton#navBtn { text-align: left; padding: 10px 12px; border: none; border-left: 2px solid transparent; border-radius: 2px; background: transparent; color: #C3CDD2; font-size: 13px; font-weight: 500; }
QPushButton#navBtn:hover { background-color: #2A353C; color: #FFFFFF; }
QPushButton#navBtn:checked { background-color: #143E52; color: #FFFFFF; border-left: 2px solid #4FB0D9; }
QLabel#pageTitle { font-size: 24px; font-weight: 700; color: #182126; letter-spacing: -0.2px; }
QLabel#pageSubtitle { font-size: 14px; color: #5C686E; }
QLabel#sectionTitle { font-size: 12px; font-weight: 700; color: #334047; letter-spacing: 0.8px; }
QLabel#sectionDesc { color: #657178; }
QFrame#card, QFrame#actionTile, QFrame#miniStat { background-color: #FAFBFA; border: 1px solid #C9D0D3; border-radius: 3px; }
QFrame#actionTile:hover { border-color: #137EAE; background-color: #F1F8FA; }
QLabel#tileIconBubble { background-color: #E5F1F5; color: #176E95; border: 1px solid #B9D6E2; border-radius: 3px; min-width: 62px; max-width: 62px; min-height: 62px; max-height: 62px; font-size: 18px; }
QLabel#tileIconBubble[tone="deep"] { background-color: #25353D; color: #FFFFFF; }
QLabel#tileTitle { font-size: 15px; color: #20282D; }
QLabel#tileSubtitle, QLabel#statLabel { color: #657178; }
QLabel#statValue { font-size: 22px; color: #20282D; }
QLabel#statMeta { color: #316B49; }
QPushButton { border-radius: 3px; padding: 7px 13px; font-size: 13px; font-weight: 600; min-height: 20px; }
QPushButton[variant="primary"] { background-color: #137EAE; color: white; border: 1px solid #0E6E99; }
QPushButton[variant="primary"]:hover { background-color: #0F709C; }
QPushButton[variant="primary"]:pressed { background-color: #0A5E84; }
QPushButton[variant="primary"]:disabled { background-color: #D6DCDE; border-color: #C4CBCE; color: #7B858A; }
QPushButton[variant="secondary"] { background-color: #FAFBFA; color: #303B41; border: 1px solid #AEB8BC; }
QPushButton[variant="secondary"]:hover { background-color: #EEF2F3; border-color: #75858D; }
QPushButton[variant="ghost"] { background: transparent; color: #126F9A; border: 1px solid transparent; }
QPushButton[variant="ghost"]:hover { background-color: #E8F2F6; }
QPushButton[variant="approve"] { background-color: #EDF7EF; color: #23673B; border: 1px solid #6FAD80; }
QPushButton[variant="approve"]:hover { background-color: #DFF0E4; }
QPushButton[variant="warning"] { background-color: #FFF7E4; color: #815500; border: 1px solid #D7A538; }
QPushButton[variant="warning"]:hover { background-color: #FFEDC0; }
QPushButton[variant="danger"] { background-color: #FFF5F3; color: #9A3C2D; border: 1px solid #D78B7E; }
QPushButton[variant="danger"]:hover { background-color: #FCE6E2; }
QLineEdit, QSpinBox, QComboBox, QListWidget, QTextEdit { background-color: #FFFFFF; border: 1px solid #B8C1C5; border-radius: 3px; padding: 6px 8px; color: #20282D; selection-background-color: #B9DCEB; }
QPushButton:focus, QLineEdit:focus, QSpinBox:focus, QComboBox:focus, QListWidget:focus, QTableWidget:focus { border: 2px solid #137EAE; }
QListWidget::item { padding: 8px; border-radius: 2px; }
QListWidget::item:selected { background-color: #DCECF3; color: #143E52; }
QTextEdit#terminal { background-color: #182126; color: #DDE5E8; border: 1px solid #11181C; border-radius: 3px; padding: 10px; }
QTableWidget { background-color: #FAFBFA; alternate-background-color: #F2F4F3; border: 1px solid #C4CCCF; gridline-color: #D8DDDF; selection-background-color: #C9E3EE; selection-color: #182126; }
QHeaderView::section { background-color: #E8ECEE; color: #3E4A50; font-weight: 700; font-size: 11px; padding: 8px 7px; border: none; border-right: 1px solid #CAD1D4; border-bottom: 1px solid #AEB8BC; }
QTableWidget::item { padding: 6px; }
QProgressBar { background-color: #D9DFE1; border: 1px solid #C0C8CB; border-radius: 2px; min-height: 7px; max-height: 7px; }
QProgressBar::chunk { background-color: #137EAE; border-radius: 1px; }
QLabel#statusChip { background-color: #E8ECEE; color: #4F5C62; padding: 4px 9px; border: 1px solid #CAD1D4; border-radius: 2px; font-size: 11px; font-weight: 700; letter-spacing: 0.6px; }
QLabel#statusChip[state="running"] { background-color: #DCEEF5; color: #0E6289; border-color: #8AC0D7; }
QLabel#statusChip[state="complete"] { background-color: #E7F4EA; color: #27683D; border-color: #9BC2A6; }
QLabel#statusChip[state="blocked"] { background-color: #FFF3D3; color: #795300; border-color: #D7B45F; }
QFrame#workflowStrip { background-color: #F9FAF9; border-bottom: 1px solid #AEB8BC; }
QFrame#drawingRegister { background-color: #EEF2F3; border-bottom: 1px solid #BFC7CA; }
QLabel#workbenchBrand { font-size: 14px; font-weight: 800; color: #1E292F; letter-spacing: 0.7px; }
QLabel#workflowStage { padding: 7px 13px; border: 1px solid #CDD4D7; border-right: none; color: #667278; font-size: 11px; font-weight: 700; }
QLabel#workflowStage[state="active"] { background-color: #137EAE; color: white; border-color: #137EAE; }
QLabel#workflowStage[state="complete"] { background-color: #EDF4F1; color: #316B49; }
QLabel#projectReference { color: #4E5B62; font-family: "Cascadia Mono", "Consolas", monospace; }
QFrame#panelNavigator, QFrame#decisionRail { background-color: #F5F7F6; }
QFrame#panelNavigator { border-right: 1px solid #BFC7CA; }
QFrame#decisionRail { border-left: 1px solid #BFC7CA; }
QLabel#railTitle, QLabel#railSection { color: #46535A; font-size: 11px; font-weight: 800; letter-spacing: 0.8px; }
QLabel#railSection { border-top: 1px solid #C7CED1; padding-top: 9px; }
QLabel#railSummary { color: #5C686E; font-size: 11px; }
QLabel#decisionTitle { font-size: 15px; font-weight: 700; color: #182126; }
QLabel#decisionMeta { color: #47545B; font-family: "Cascadia Mono", "Consolas", monospace; font-size: 11px; }
QLabel#evidencePreview { background-color: #FFFFFF; border: 1px solid #BEC7CA; color: #6C777C; }
QTextEdit#evidenceText { font-family: "Cascadia Mono", "Consolas", monospace; font-size: 10px; }
QLabel#conflictText { color: #795300; background-color: #FFF8E8; border: 1px solid #E1C77D; padding: 7px; }
QFrame#tableToolbar { background-color: #F5F7F6; border-bottom: 1px solid #BFC7CA; }
QLabel#tableTitle { font-size: 12px; font-weight: 800; color: #334047; letter-spacing: 0.7px; }
QTabWidget#reviewTabs::pane { border: 1px solid #BFC7CA; background: #FAFBFA; }
QTabBar::tab { background: #E8ECEE; border: 1px solid #C4CCCF; border-bottom: none; padding: 7px 13px; }
QTabBar::tab:selected { background: #FAFBFA; color: #126F9A; }
QFrame#reviewFooter { background-color: #F9FAF9; border-top: 1px solid #AEB8BC; }
QLabel#gateLabel { font-size: 11px; font-weight: 800; letter-spacing: 0.5px; }
QLabel#gateLabel[state="ready"] { color: #27683D; }
QLabel#gateLabel[state="blocked"] { color: #795300; }
QLabel#saveState { color: #5F6C72; font-size: 11px; }
QLabel#saveState[state="dirty"] { color: #8A5B00; }
"""


@dataclass
class AppConfig:
    gemini_api_key: str = ""
    gemini_model: str = DEFAULT_GEMINI_MODEL
    db_path: str = str(PROJECT_ROOT / "pricing.db")
    python_executable: str = sys.executable
    default_year: int = 2026
    default_half: str = "H1"
    default_pricing_mode: str = "total"


@dataclass
class JobRecord:
    finished_at: str
    project_name: str
    brands: list[str]
    status: str
    output_dir: str
    error: str = ""


@dataclass
class DbColumnInfo:
    name: str
    declared_type: str
    not_null: bool
    pk_pos: int


def _load_json(path: Path, default: object) -> object:
    if not path.exists():
        return default
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default


def _save_json(path: Path, payload: object) -> None:
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def load_config() -> AppConfig:
    cfg = AppConfig()
    data = _load_json(CONFIG_PATH, {})
    if isinstance(data, dict):
        for key, value in data.items():
            if hasattr(cfg, key):
                setattr(cfg, key, value)
    if not str(cfg.gemini_model or "").strip():
        cfg.gemini_model = DEFAULT_GEMINI_MODEL
    return cfg


def save_config(cfg: AppConfig) -> None:
    _save_json(CONFIG_PATH, asdict(cfg))


def load_history() -> list[JobRecord]:
    rows = _load_json(HISTORY_PATH, [])
    out: list[JobRecord] = []
    if isinstance(rows, list):
        for row in rows:
            if isinstance(row, dict):
                try:
                    out.append(JobRecord(**row))
                except TypeError:
                    pass
    return out


def save_history(rows: list[JobRecord]) -> None:
    _save_json(HISTORY_PATH, [asdict(r) for r in rows])


def sanitize_filename(value: str) -> str:
    clean = re.sub(r"[\\/:*?\"<>|]", "_", value.strip())
    clean = re.sub(r"\s+", "_", clean)
    return clean or "project"


def mask_api_key(value: str) -> str:
    if not value:
        return ""
    if len(value) < 8:
        return "****"
    return f"{value[:4]}...{value[-4:]}"


def fetch_brands(db_path: str) -> list[str]:
    db = Path(db_path)
    if not db.exists():
        return []
    try:
        conn = sqlite3.connect(str(db))
        rows = conn.execute("SELECT name FROM manufacturers ORDER BY name").fetchall()
        conn.close()
        return sorted({str(r[0]).strip().upper() for r in rows if r and r[0]})
    except Exception:
        return []


def resolve_oda_converter_path() -> str:
    env_path = (os.getenv("ODA_FILE_CONVERTER") or "").strip()
    candidates: list[Path] = []
    if env_path:
        candidates.append(Path(env_path).expanduser())

    candidates.extend(
        [
            PROJECT_ROOT / "ODA" / "ODAFileConverter.exe",
            SOURCE_ROOT / "ODA" / "ODAFileConverter.exe",
            BUNDLE_ROOT / "ODA" / "ODAFileConverter.exe",
        ]
    )
    for path in candidates:
        if path.exists():
            return str(path)
    return ""


def quote_ident(value: str) -> str:
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", value or ""):
        raise ValueError(f"Invalid SQL identifier: {value!r}")
    return f'"{value}"'


def list_sqlite_tables(conn: sqlite3.Connection) -> list[str]:
    rows = conn.execute(
        "SELECT name FROM sqlite_master "
        "WHERE type='table' AND name NOT LIKE 'sqlite_%' "
        "ORDER BY name"
    ).fetchall()
    return [str(r[0]) for r in rows if r and r[0]]


def get_table_columns(conn: sqlite3.Connection, table_name: str) -> list[DbColumnInfo]:
    pragma_rows = conn.execute(f"PRAGMA table_info({quote_ident(table_name)})").fetchall()
    cols: list[DbColumnInfo] = []
    for row in pragma_rows:
        cols.append(
            DbColumnInfo(
                name=str(row[1]),
                declared_type=str(row[2] or ""),
                not_null=bool(row[3]),
                pk_pos=int(row[5] or 0),
            )
        )
    return cols


def test_gemini(api_key: str, model: str) -> tuple[bool, str]:
    if not api_key:
        return False, "Gemini API key is empty."
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={api_key}"
    payload = json.dumps({"contents": [{"parts": [{"text": "Reply with OK"}]}]}).encode("utf-8")
    req = urllib.request.Request(url, data=payload, headers={"Content-Type": "application/json"}, method="POST")
    try:
        with urllib.request.urlopen(req, timeout=15) as resp:
            body = json.loads(resp.read().decode("utf-8", errors="replace"))
        text = body["candidates"][0]["content"]["parts"][0]["text"]
        return True, f"Connected. Model response: {text[:80]}"
    except urllib.error.HTTPError as exc:
        return False, f"HTTP {exc.code}: {exc.reason}"
    except Exception as exc:
        return False, f"Connection failed: {exc}"


def set_button_variant(button: QPushButton, variant: str) -> None:
    button.setProperty("variant", variant)
    button.style().unpolish(button)
    button.style().polish(button)


class ActionTile(QFrame):
    clicked = Signal()

    def __init__(self, icon_glyph: str, title: str, subtitle: str = "", tone: str = "default", parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setObjectName("actionTile")
        self.setCursor(Qt.PointingHandCursor)
        self.setMinimumHeight(220)
        self.setFocusPolicy(Qt.StrongFocus)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 24, 24, 24)
        layout.setSpacing(10)
        layout.setAlignment(Qt.AlignCenter)

        # Icon
        icon_lbl = QLabel(icon_glyph)
        icon_lbl.setObjectName("tileIconBubble")
        icon_lbl.setProperty("tone", tone)
        icon_lbl.setAlignment(Qt.AlignCenter)
        icon_lbl.setFont(QFont("Segoe UI Symbol", 32))
        
        # Text
        title_label = QLabel(title)
        title_label.setObjectName("tileTitle")
        title_label.setAlignment(Qt.AlignCenter)
        subtitle_label = QLabel(subtitle)
        subtitle_label.setObjectName("tileSubtitle")
        subtitle_label.setAlignment(Qt.AlignCenter)
        subtitle_label.setWordWrap(True)

        layout.addWidget(icon_lbl, 0, Qt.AlignHCenter)
        layout.addSpacing(6)
        layout.addWidget(title_label)
        if subtitle:
            layout.addWidget(subtitle_label)

    def mousePressEvent(self, event) -> None:  # type: ignore[override]
        if event.button() == Qt.LeftButton:
            self.clicked.emit()
        super().mousePressEvent(event)


class MiniStatCard(QFrame):
    def __init__(self, label: str, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setObjectName("miniStat")

        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(4)

        self.value_label = QLabel("--")
        self.value_label.setObjectName("statValue")
        self.label = QLabel(label)
        self.label.setObjectName("statLabel")
        self.meta = QLabel("")
        self.meta.setObjectName("statMeta")

        layout.addWidget(self.label)
        layout.addWidget(self.value_label)
        layout.addWidget(self.meta)

    def set_values(self, value: str, meta: str) -> None:
        self.value_label.setText(value)
        self.meta.setText(meta)


class WideCellEditorDelegate(QStyledItemDelegate):
    def __init__(self, extra_width: int = 180, min_width: int = 260, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.extra_width = extra_width
        self.min_width = min_width

    def updateEditorGeometry(self, editor, option, index) -> None:  # type: ignore[override]
        rect = QRect(option.rect)
        target = max(self.min_width, rect.width() + self.extra_width)
        viewport = editor.parentWidget()
        if viewport is not None:
            available = max(80, viewport.width() - rect.left() - 6)
            target = min(target, available)
        rect.setWidth(max(rect.width(), target))
        editor.setGeometry(rect)


class SettingsDialog(QDialog):
    def __init__(self, cfg: AppConfig, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setWindowTitle("Preferences")
        self.resize(650, 480)

        self.api_key = QLineEdit(cfg.gemini_api_key)
        self.api_key.setEchoMode(QLineEdit.Password)
        self.model = QLineEdit(cfg.gemini_model)
        self.db = QLineEdit(cfg.db_path)
        self.py = QLineEdit(cfg.python_executable)
        self.default_year = QSpinBox()
        self.default_year.setRange(2024, 2050)
        self.default_year.setValue(cfg.default_year)
        self.default_half = QComboBox()
        self.default_half.addItems(["H1", "H2"])
        self.default_half.setCurrentText(cfg.default_half)
        self.default_mode = QComboBox()
        self.default_mode.addItems(["total", "itemized"])
        self.default_mode.setCurrentText(cfg.default_pricing_mode)

        form = QFormLayout()
        form.setHorizontalSpacing(24)
        form.setVerticalSpacing(16)
        
        form.addRow("Gemini API Key", self._row_with_buttons(self.api_key, [self._btn("Toggle", self._toggle_key, "secondary")]))
        form.addRow("Gemini Model", self.model)
        form.addRow("Pricing DB", self._row_with_buttons(self.db, [self._btn("Browse", lambda: self._pick_file(self.db, "SQLite DB (*.db);;All Files (*)"), "secondary")]))
        if not IS_FROZEN:
            form.addRow("Python Executable", self._row_with_buttons(self.py, [self._btn("Browse", lambda: self._pick_file(self.py, "Executable (*.exe);;All Files (*)"), "secondary")]))
        form.addRow("Default Year", self.default_year)
        form.addRow("Default Half", self.default_half)
        form.addRow("Default Mode", self.default_mode)

        test_btn = QPushButton("Test API Connection")
        set_button_variant(test_btn, "secondary")
        test_btn.clicked.connect(self._test)

        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        
        save_btn = buttons.button(QDialogButtonBox.Save)
        cancel_btn = buttons.button(QDialogButtonBox.Cancel)
        if save_btn: set_button_variant(save_btn, "primary")
        if cancel_btn: set_button_variant(cancel_btn, "ghost")

        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 24, 24, 24)
        layout.setSpacing(24)
        
        header = QLabel("Application Preferences")
        header.setObjectName("pageTitle")
        
        container = QFrame()
        container.setObjectName("card")
        body = QVBoxLayout(container)
        body.setContentsMargins(24, 24, 24, 24)
        body.addLayout(form)
        body.addSpacing(12)
        body.addWidget(test_btn, 0, Qt.AlignLeft)

        layout.addWidget(header)
        layout.addWidget(container)
        layout.addStretch(1)
        layout.addWidget(buttons)

    def _btn(self, text: str, fn, variant: str) -> QPushButton:
        btn = QPushButton(text)
        btn.clicked.connect(fn)
        set_button_variant(btn, variant)
        return btn

    def _row_with_buttons(self, line: QLineEdit, buttons: list[QPushButton]) -> QWidget:
        row = QWidget()
        lay = QHBoxLayout(row)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(8)
        lay.addWidget(line, 1)
        for btn in buttons:
            lay.addWidget(btn)
        return row

    def _pick_file(self, target: QLineEdit, flt: str) -> None:
        path, _ = QFileDialog.getOpenFileName(self, "Select File", str(PROJECT_ROOT), flt)
        if path:
            target.setText(path)

    def _toggle_key(self) -> None:
        mode = QLineEdit.Normal if self.api_key.echoMode() == QLineEdit.Password else QLineEdit.Password
        self.api_key.setEchoMode(mode)

    def _test(self) -> None:
        ok, message = test_gemini(self.api_key.text().strip(), self.model.text().strip() or DEFAULT_GEMINI_MODEL)
        if ok:
            QMessageBox.information(self, "Success", message)
        else:
            QMessageBox.warning(self, "Connection Failed", message)

    def to_config(self) -> AppConfig:
        python_exec = sys.executable if IS_FROZEN else (self.py.text().strip() or sys.executable)
        return AppConfig(
            gemini_api_key=self.api_key.text().strip(),
            gemini_model=self.model.text().strip() or DEFAULT_GEMINI_MODEL,
            db_path=self.db.text().strip(),
            python_executable=python_exec,
            default_year=int(self.default_year.value()),
            default_half=self.default_half.currentText(),
            default_pricing_mode=self.default_mode.currentText(),
        )


class ImportExcelDialog(QDialog):
    def __init__(self, cfg: AppConfig, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setWindowTitle("Import Excel to Pricing DB")
        self.resize(760, 420)

        self.excel_path = QLineEdit()
        self.db_path = QLineEdit(cfg.db_path or str(PROJECT_ROOT / "pricing.db"))
        self.manufacturer = QLineEdit()
        self.year = QSpinBox()
        self.year.setRange(2024, 2050)
        self.year.setValue(int(cfg.default_year))
        self.half = QComboBox()
        self.half.addItems(["H1", "H2"])
        self.half.setCurrentText(cfg.default_half)
        self.currency = QComboBox()
        self.currency.addItems(["EUR", "USD"])
        self.currency.setCurrentText("EUR")
        self.force = QCheckBox("Force import even if AI validation warns")

        form = QFormLayout()
        form.setHorizontalSpacing(24)
        form.setVerticalSpacing(16)
        form.addRow(
            "Excel File",
            self._row_with_buttons(
                self.excel_path,
                [self._btn("Browse", self._pick_excel, "secondary")],
            ),
        )
        form.addRow(
            "Output DB",
            self._row_with_buttons(
                self.db_path,
                [self._btn("Browse", self._pick_db, "secondary")],
            ),
        )
        form.addRow("Manufacturer", self.manufacturer)
        form.addRow("Year", self.year)
        form.addRow("Half", self.half)
        form.addRow("Currency", self.currency)
        form.addRow("Options", self.force)

        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        ok_btn = buttons.button(QDialogButtonBox.Ok)
        cancel_btn = buttons.button(QDialogButtonBox.Cancel)
        if ok_btn is not None:
            ok_btn.setText("Import")
            set_button_variant(ok_btn, "primary")
        if cancel_btn is not None:
            set_button_variant(cancel_btn, "ghost")
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 24, 24, 24)
        layout.setSpacing(18)

        title = QLabel("Import Manufacturer Price List")
        title.setObjectName("pageTitle")
        subtitle = QLabel(
            "Select an Excel file and import it into a SQLite DB. "
            "The app will automatically switch Preferences to the imported DB."
        )
        subtitle.setObjectName("pageSubtitle")
        subtitle.setWordWrap(True)

        card = QFrame()
        card.setObjectName("card")
        card_layout = QVBoxLayout(card)
        card_layout.setContentsMargins(24, 24, 24, 24)
        card_layout.addLayout(form)

        layout.addWidget(title)
        layout.addWidget(subtitle)
        layout.addWidget(card)
        layout.addStretch(1)
        layout.addWidget(buttons)

    def _btn(self, text: str, fn, variant: str) -> QPushButton:
        btn = QPushButton(text)
        btn.clicked.connect(fn)
        set_button_variant(btn, variant)
        return btn

    def _row_with_buttons(self, line: QLineEdit, buttons: list[QPushButton]) -> QWidget:
        row = QWidget()
        lay = QHBoxLayout(row)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(8)
        lay.addWidget(line, 1)
        for btn in buttons:
            lay.addWidget(btn)
        return row

    def _pick_excel(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Select Excel file",
            str(PROJECT_ROOT),
            "Excel Files (*.xlsx *.xlsm *.xls);;All Files (*)",
        )
        if not path:
            return
        self.excel_path.setText(path)
        xlsx = Path(path)
        if not self.db_path.text().strip():
            self.db_path.setText(str(xlsx.with_suffix(".db")))
        if not self.manufacturer.text().strip():
            stem_upper = xlsx.stem.upper()
            for brand in DEFAULT_BRANDS:
                if brand in stem_upper:
                    self.manufacturer.setText(brand)
                    break

    def _pick_db(self) -> None:
        start = self.db_path.text().strip() or str(PROJECT_ROOT / "pricing_import.db")
        path, _ = QFileDialog.getSaveFileName(
            self,
            "Choose output DB file",
            start,
            "SQLite DB (*.db);;All Files (*)",
        )
        if path:
            self.db_path.setText(path)

    def payload(self) -> dict:
        return {
            "excel_file": self.excel_path.text().strip(),
            "db_path": self.db_path.text().strip(),
            "manufacturer": self.manufacturer.text().strip().upper(),
            "year": int(self.year.value()),
            "half": self.half.currentText(),
            "currency": self.currency.currentText(),
            "force": bool(self.force.isChecked()),
        }

    def accept(self) -> None:  # type: ignore[override]
        data = self.payload()
        excel = Path(data["excel_file"])
        if not data["excel_file"]:
            QMessageBox.warning(self, "Missing Excel File", "Please select an Excel file to import.")
            return
        if not excel.exists():
            QMessageBox.warning(self, "Excel Not Found", f"File not found:\n{excel}")
            return
        if not data["manufacturer"]:
            QMessageBox.warning(self, "Missing Manufacturer", "Please enter a manufacturer name.")
            return
        if not data["db_path"]:
            QMessageBox.warning(self, "Missing Output DB", "Please choose an output DB path.")
            return
        super().accept()


class ExcelImportWorker(QObject):
    finished = Signal(dict)
    failed = Signal(str)

    def __init__(
        self,
        *,
        python_executable: str,
        script_path: Path,
        script_args: list[str],
        env: dict[str, str],
        db_path: Path,
        manufacturer: str,
        year: int,
        half: str,
    ) -> None:
        super().__init__()
        self.python_executable = python_executable
        self.script_path = script_path
        self.script_args = script_args
        self.env = env
        self.db_path = db_path
        self.manufacturer = manufacturer
        self.year = year
        self.half = half

    def run(self) -> None:
        try:
            code, output = self._run_python_script()
            if code != 0:
                tail = "\n".join((output or "").strip().splitlines()[-20:])
                raise RuntimeError(f"Import process exited with code {code}.\n{tail}")

            summary = self._import_result_summary()
            if summary["edition_products"] <= 0:
                tail = "\n".join((output or "").strip().splitlines()[-20:])
                raise RuntimeError(
                    "Import completed but no products were written for the selected "
                    f"manufacturer/year/half.\n{tail}"
                )

            self.finished.emit(
                {
                    "db_path": str(self.db_path),
                    "summary": summary,
                    "output": output,
                }
            )
        except Exception as exc:
            self.failed.emit(str(exc))

    def _run_python_script(self) -> tuple[int, str]:
        if IS_FROZEN:
            return self._run_python_script_embedded()

        cmd = [self.python_executable, str(self.script_path), *self.script_args]
        proc = subprocess.run(
            cmd,
            cwd=str(PROJECT_ROOT),
            env=self.env,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        return proc.returncode, proc.stdout or ""

    def _run_python_script_embedded(self) -> tuple[int, str]:
        lines: list[str] = []

        class _Stream(io.TextIOBase):
            def __init__(self, sink: list[str]) -> None:
                self._buf = ""
                self._sink = sink

            def write(self, s: str) -> int:
                self._buf += s
                while "\n" in self._buf:
                    line, self._buf = self._buf.split("\n", 1)
                    self._sink.append(line.rstrip("\r"))
                return len(s)

            def flush(self) -> None:
                if self._buf:
                    self._sink.append(self._buf.rstrip("\r"))
                    self._buf = ""

        overrides = {}
        for key in ("GEMINI_API_KEY", "GEMINI_API_MODEL", "PYTHONUTF8", "PYTHONIOENCODING"):
            if key in self.env:
                overrides[key] = self.env[key]

        prev_env = {k: os.environ.get(k) for k in overrides}
        for k, v in overrides.items():
            os.environ[k] = v

        old_argv = sys.argv[:]
        old_cwd = Path.cwd()
        stream = _Stream(lines)
        code = 0

        try:
            os.chdir(PROJECT_ROOT)
            sys.argv = [str(self.script_path), *self.script_args]
            with redirect_stdout(stream), redirect_stderr(stream):
                try:
                    runpy.run_path(str(self.script_path), run_name="__main__")
                except SystemExit as exc:
                    code = exc.code if isinstance(exc.code, int) else (0 if exc.code in (None, "") else 1)
        except Exception as exc:
            code = 1
            lines.append(str(exc))
        finally:
            stream.flush()
            os.chdir(old_cwd)
            sys.argv = old_argv
            for k, old in prev_env.items():
                if old is None:
                    os.environ.pop(k, None)
                else:
                    os.environ[k] = old

        return code, "\n".join(lines)

    def _import_result_summary(self) -> dict[str, int]:
        conn = sqlite3.connect(str(self.db_path))
        try:
            total_products = int(conn.execute("SELECT COUNT(*) FROM products").fetchone()[0])
            edition_products = int(
                conn.execute(
                    """
                    SELECT COUNT(*)
                    FROM products p
                    JOIN ppl_editions pe ON pe.id = p.ppl_edition_id
                    JOIN manufacturers m ON m.id = pe.manufacturer_id
                    WHERE UPPER(m.name) = ? AND pe.year = ? AND pe.half = ?
                    """,
                    (self.manufacturer.upper(), int(self.year), self.half),
                ).fetchone()[0]
            )
        finally:
            conn.close()
        return {
            "total_products": total_products,
            "edition_products": edition_products,
        }


class ProjectWorker(QObject):
    log = Signal(str)
    step = Signal(int, int, str)
    finished = Signal(dict)
    failed = Signal(str, dict)
    cancelled = Signal(dict)

    def __init__(self, cfg: AppConfig, project_name: str, project_ref: str, drawing_paths: list[Path], output_dir: Path, brands: list[str], year: int, half: str, pricing_mode: str) -> None:
        super().__init__()
        self.cfg = cfg
        self.project_name = project_name
        self.project_ref = project_ref
        self.drawing_paths = drawing_paths
        self.output_dir = output_dir
        self.brands = [b.upper() for b in brands]
        self.year = year
        self.half = half
        self.pricing_mode = pricing_mode

    def run(self) -> None:
        ctx = {
            "project_name": self.project_name,
            "project_ref": self.project_ref,
            "brands": self.brands,
            "output_dir": str(self.output_dir),
            "raw_output": "",
            "offers": [],
            "finished_at": "",
        }
        try:
            run_dir = self.output_dir / sanitize_filename(self.project_name)
            run_dir.mkdir(parents=True, exist_ok=True)
            ctx["output_dir"] = str(run_dir)
            api_key, model = self._resolve_gemini_runtime_config()
            oda_path = resolve_oda_converter_path()
            self.step.emit(1, 1, "Extracting components into canonical project JSON")
            editions = {
                brand: {"year": self.year, "half": self.half}
                for brand in self.brands
            }
            pipeline = ProjectPipeline(
                database=self.cfg.db_path,
                api_key=api_key,
                oda_path=oda_path,
                progress=self.log.emit,
                should_cancel=lambda: QThread.currentThread().isInterruptionRequested(),
            )
            result = pipeline.run(
                project_dir=run_dir,
                project_name=self.project_name,
                project_reference=self.project_ref,
                drawings=self.drawing_paths,
                default_manufacturer=self.brands[0],
                editions=editions,
                pricing_mode=self.pricing_mode,
                price_after_extraction=False,
            )
            ctx["raw_output"] = result["project_path"]
            ctx["pricing_output"] = ""

            ctx["finished_at"] = datetime.now().replace(microsecond=0).isoformat()
            self.finished.emit(ctx)
        except InterruptedError:
            ctx["finished_at"] = datetime.now().replace(microsecond=0).isoformat()
            self.cancelled.emit(ctx)
        except Exception as exc:
            ctx["finished_at"] = datetime.now().replace(microsecond=0).isoformat()
            self.failed.emit(str(exc), ctx)

    def _resolve_gemini_runtime_config(self) -> tuple[str, str]:
        cfg_model = (self.cfg.gemini_model or "").strip()
        model = cfg_model or DEFAULT_GEMINI_MODEL
        key = (self.cfg.gemini_api_key or "").strip()

        legacy_env = PROJECT_ROOT / ".venv" / "key.env"
        if legacy_env.exists():
            for line in legacy_env.read_text(encoding="utf-8", errors="replace").splitlines():
                raw = line.strip()
                if raw.startswith("GEMINI_API_KEY=") and not key:
                    key = raw.split("=", 1)[1].strip()
                if raw.startswith("GEMINI_API_MODEL="):
                    file_model = raw.split("=", 1)[1].strip()
                    if file_model and not cfg_model:
                        model = file_model

        if not key:
            raise RuntimeError(
                "Gemini API key is required. Set it in Preferences or add GEMINI_API_KEY to .venv/key.env."
            )

        # Keep legacy script expectations satisfied without modifying script code.
        legacy_env.parent.mkdir(parents=True, exist_ok=True)
        legacy_env.write_text(
            f"GEMINI_API_KEY={key}\nGEMINI_API_MODEL={model}\n"
            "GEMINI_FALLBACK_MODEL=gemini-3.1-flash-lite\n",
            encoding="utf-8",
        )
        return key, model

    def _prepare_pricing_input(self, raw_path: Path, run_dir: Path) -> Path:
        try:
            from openpyxl import Workbook, load_workbook
        except Exception as exc:
            raise RuntimeError(f"openpyxl is required for pricing input preparation: {exc}") from exc

        normalized = run_dir / f"{raw_path.stem}_match_offer.xlsx"
        wb = load_workbook(str(raw_path), data_only=True)
        ws = wb[wb.sheetnames[0]]

        def norm(value) -> str:
            return str(value).strip().lower() if value is not None else ""

        def to_int(value, default: int) -> int:
            try:
                return int(float(value))
            except Exception:
                return default

        def find_header_row(second_col_text: str) -> int | None:
            target = second_col_text.strip().lower()
            for row in range(1, ws.max_row + 1):
                if norm(ws.cell(row, 1).value) == "#" and norm(ws.cell(row, 2).value) == target:
                    return row
            return None

        components_header = find_header_row("category")
        terminals_header = find_header_row("terminal type")
        if components_header is None or terminals_header is None:
            wb.close()
            raise RuntimeError("Could not find expected components/terminals headers in extraction workbook.")

        def map_cols(header_row: int) -> dict[str, int]:
            out: dict[str, int] = {}
            for col in range(1, ws.max_column + 1):
                key = norm(ws.cell(header_row, col).value)
                if key:
                    out[key] = col
            return out

        cm = map_cols(components_header)
        tm = map_cols(terminals_header)

        def getc(row: int, key: str, default=None):
            col = cm.get(key)
            if col is None and key == "b.c.":
                col = cm.get("b.c")
            return ws.cell(row, col).value if col else default

        components: list[dict] = []
        for row in range(components_header + 1, terminals_header):
            idx = ws.cell(row, 1).value
            if not isinstance(idx, (int, float)):
                continue
            raw_qty = to_int(getc(row, "qty", 1), 1)
            qty = raw_qty if raw_qty > 0 else 1
            terminal_flag = str(getc(row, "terminal?", "") or "").strip()
            has_terminal = terminal_flag in {"\u2713", "\u2714", "yes", "Yes", "true", "True", "1"}
            components.append(
                {
                    "category": str(getc(row, "category", "OTHER") or "OTHER").strip(),
                    "label": str(getc(row, "label", "-") or "-").strip(),
                    "poles": str(getc(row, "poles", "-") or "-").strip(),
                    "amps": str(getc(row, "amps", "-") or "-").strip(),
                    "curve": str(getc(row, "curve", "-") or "-").strip(),
                    "breaking_capacity": str(getc(row, "b.c.", "-") or "-").strip(),
                    "sensitivity": str(getc(row, "sensitivity", "-") or "-").strip(),
                    "spd_type": str(getc(row, "spd type", "-") or "-").strip(),
                    "qty": qty,
                    "has_terminal": has_terminal,
                }
            )

        terminals: list[dict] = []
        name_col = tm.get("terminal type", 2)
        qty_col = tm.get("qty", 5)
        for row in range(terminals_header + 1, ws.max_row + 1):
            idx = ws.cell(row, 1).value
            if not isinstance(idx, (int, float)):
                continue
            name = ws.cell(row, name_col).value
            qty = ws.cell(row, qty_col).value
            if not name:
                continue
            q = to_int(qty, 0)
            if q <= 0:
                continue
            terminals.append({"description": str(name).strip(), "qty": q})

        panel_name = ws.title.strip() or raw_path.stem
        title_cell = ws.cell(1, 1).value
        if isinstance(title_cell, str) and "\u00b7" in title_cell:
            extracted = title_cell.split("\u00b7", 1)[0].strip()
            if extracted:
                panel_name = extracted
        wb.close()

        out = Workbook()
        enc = out.active
        enc.title = "Enclosure"
        enc["A4"] = "Panel Reference"
        enc["B4"] = panel_name

        comp_ws = out.create_sheet("Components")
        comp_headers = ["#", "Category", "Label", "Poles", "Amps", "Curve", "B.C.", "Sensitivity", "SPD Type", "Qty", "Terminal?"]
        for i, header in enumerate(comp_headers, start=1):
            comp_ws.cell(row=3, column=i, value=header)
        for i, item in enumerate(components, start=4):
            comp_ws.cell(i, 1, i - 3)
            comp_ws.cell(i, 2, item["category"])
            comp_ws.cell(i, 3, item["label"])
            comp_ws.cell(i, 4, item["poles"])
            comp_ws.cell(i, 5, item["amps"])
            comp_ws.cell(i, 6, item["curve"])
            comp_ws.cell(i, 7, item["breaking_capacity"])
            comp_ws.cell(i, 8, item["sensitivity"])
            comp_ws.cell(i, 9, item["spd_type"])
            comp_ws.cell(i, 10, item["qty"])
            comp_ws.cell(i, 11, "\u2713" if item["has_terminal"] else "-")

        term_ws = out.create_sheet("Terminal Schedule")
        term_headers = ["#", "Terminal Type", "Qty"]
        for i, header in enumerate(term_headers, start=1):
            term_ws.cell(row=3, column=i, value=header)
        for i, item in enumerate(terminals, start=4):
            term_ws.cell(i, 1, i - 3)
            term_ws.cell(i, 2, item["description"])
            term_ws.cell(i, 3, item["qty"])

        normalized.parent.mkdir(parents=True, exist_ok=True)
        out.save(str(normalized))
        self.log.emit(f"Prepared pricing input -> {normalized}")
        self.log.emit(f"Components: {len(components)} | Terminals: {len(terminals)} | Panel: {panel_name}")
        return normalized

    def _run_embedded_script(self, script_path: Path, script_args: list[str], env: dict[str, str]) -> None:
        tail: list[str] = []

        def emit_line(line: str) -> None:
            text = line.rstrip("\r")
            if text:
                self.log.emit(text)
                tail.append(text)
                if len(tail) > 20:
                    tail.pop(0)

        class _Stream(io.TextIOBase):
            def __init__(self, on_line) -> None:
                self._buf = ""
                self._on_line = on_line

            def write(self, s: str) -> int:
                self._buf += s
                while "\n" in self._buf:
                    line, self._buf = self._buf.split("\n", 1)
                    self._on_line(line)
                return len(s)

            def flush(self) -> None:
                if self._buf:
                    self._on_line(self._buf)
                    self._buf = ""

        overrides = {}
        for key in ("GEMINI_API_KEY", "GEMINI_API_MODEL"):
            if key in env:
                overrides[key] = env[key]
        prev_env = {k: os.environ.get(k) for k in overrides}
        for k, v in overrides.items():
            os.environ[k] = v

        old_argv = sys.argv[:]
        old_cwd = Path.cwd()
        stream = _Stream(emit_line)
        try:
            os.chdir(PROJECT_ROOT)
            sys.argv = [str(script_path), *script_args]
            # PyInstaller can miss package attribute exports used by legacy imports
            # like `from ezdxf import recover` in embedded scripts.
            try:
                import ezdxf  # type: ignore[import-not-found]
                import ezdxf.recover as ezdxf_recover  # type: ignore[import-not-found]
                if getattr(ezdxf, "recover", None) is None:
                    setattr(ezdxf, "recover", ezdxf_recover)
            except Exception:
                pass
            with redirect_stdout(stream), redirect_stderr(stream):
                try:
                    runpy.run_path(str(script_path), run_name="__main__")
                except SystemExit as exc:
                    code = exc.code if isinstance(exc.code, int) else (0 if exc.code in (None, "") else 1)
                    if code != 0:
                        raise RuntimeError(f"Script exited with code {code}.")
        except Exception as exc:
            raise RuntimeError(f"{exc}\n" + "\n".join(tail[-8:])) from exc
        finally:
            stream.flush()
            os.chdir(old_cwd)
            sys.argv = old_argv
            for k, old in prev_env.items():
                if old is None:
                    os.environ.pop(k, None)
                else:
                    os.environ[k] = old

    def _run(self, cmd: list[str], env: dict[str, str]) -> None:
        preview = self._cmd_preview(cmd)
        self.log.emit(f"\n> {preview}\n")

        if IS_FROZEN and len(cmd) >= 3 and cmd[1].lower().endswith(".py"):
            self._run_embedded_script(Path(cmd[1]), cmd[2:], env)
            return

        proc = subprocess.Popen(
            cmd,
            cwd=str(PROJECT_ROOT),
            env=env,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        assert proc.stdout is not None
        tail: list[str] = []
        for line in proc.stdout:
            text = line.rstrip()
            self.log.emit(text)
            tail.append(text)
            if len(tail) > 20:
                tail.pop(0)
        code = proc.wait()
        if code != 0:
            raise RuntimeError(f"Process exited with code {code}.\n" + "\n".join(tail[-8:]))
    def _cmd_preview(self, cmd: list[str]) -> str:
        out: list[str] = []
        skip = False
        for i, token in enumerate(cmd):
            if skip:
                skip = False
                continue
            if token == "--api-key" and i + 1 < len(cmd):
                out.extend([token, mask_api_key(cmd[i + 1])])
                skip = True
            elif token == "-c" and i + 1 < len(cmd):
                out.extend([token, "<inline-script>"])
                skip = True
            else:
                out.append(token)
        return " ".join(out)


class HomePage(QWidget):
    new_project_requested = Signal()
    history_requested = Signal()
    pricing_db_requested = Signal()
    import_excel_requested = Signal()

    def __init__(self) -> None:
        super().__init__()
        lay = QVBoxLayout(self)
        lay.setContentsMargins(28, 28, 28, 28)
        lay.setSpacing(24)

        header_lay = QVBoxLayout()
        header_lay.setSpacing(8)
        title = QLabel("Engineering workbench")
        title.setObjectName("pageTitle")
        desc = QLabel("Register source drawings, inspect or edit the evidence-backed BOM, then price and export all included components.")
        desc.setObjectName("pageSubtitle")
        desc.setWordWrap(True)
        header_lay.addWidget(title)
        header_lay.addWidget(desc)

        stats_row = QHBoxLayout()
        stats_row.setSpacing(24)
        self.stat_db_update = MiniStatCard("Last DB Update")
        self.stat_last_success = MiniStatCard("Last Successful Project")
        self.stat_recent = MiniStatCard("Latest Pipeline Run")
        stats_row.addWidget(self.stat_db_update, 1)
        stats_row.addWidget(self.stat_last_success, 1)
        stats_row.addWidget(self.stat_recent, 1)

        actions_lay = QVBoxLayout()
        actions_lay.setSpacing(16)
        action_title = QLabel("Quick Actions")
        action_title.setObjectName("sectionTitle")
        
        tile_row = QHBoxLayout()
        tile_row.setSpacing(24)
        new_tile = ActionTile(
            "NEW",
            "New Project",
            "Start workflow",
            "default",
        )
        history_tile = ActionTile(
            "RUNS",
            "Job History",
            "View past runs",
            "deep",
        )
        pricing_db_tile = ActionTile(
            "DB",
            "Pricing DB",
            "Read/edit database",
            "default",
        )
        import_excel_tile = ActionTile(
            "XLS",
            "Import Excel",
            "Convert Excel to DB",
            "default",
        )
        tile_row.addWidget(new_tile, 1)
        tile_row.addWidget(history_tile, 1)
        tile_row.addWidget(pricing_db_tile, 1)
        tile_row.addWidget(import_excel_tile, 1)
        actions_lay.addWidget(action_title)
        actions_lay.addLayout(tile_row)

        lay.addLayout(header_lay)
        lay.addLayout(stats_row)
        lay.addSpacing(16)
        lay.addLayout(actions_lay)
        lay.addStretch(1)

        new_tile.clicked.connect(self.new_project_requested.emit)
        history_tile.clicked.connect(self.history_requested.emit)
        pricing_db_tile.clicked.connect(self.pricing_db_requested.emit)
        import_excel_tile.clicked.connect(self.import_excel_requested.emit)

    def set_overview(self, records: list[JobRecord], db_path: str = "") -> None:
        db_value = "--"
        db_meta = "Configure Pricing DB in Preferences"
        db_raw = str(db_path or "").strip()
        if db_raw:
            db_file = Path(db_raw).expanduser()
            if db_file.exists():
                try:
                    modified_dt = datetime.fromtimestamp(db_file.stat().st_mtime)
                    db_value = modified_dt.strftime("%B %d, %Y")
                    db_meta = f"{modified_dt.strftime('%I:%M %p')} · {db_file.name}"
                except Exception:
                    db_value = "Available"
                    db_meta = db_file.name
            else:
                db_value = "Missing"
                db_meta = db_file.name

        success_value = "--"
        success_meta = "No successful runs yet"
        last_success = next((r for r in reversed(records) if r.status == "success"), None)
        if last_success is not None:
            success_value = last_success.project_name or "(Unnamed project)"
            try:
                success_dt = datetime.fromisoformat(last_success.finished_at)
                success_meta = success_dt.strftime("%B %d, %Y · %I:%M %p")
            except Exception:
                success_meta = last_success.finished_at[:16] or "time unavailable"

            if last_success.brands:
                success_meta += f" · {len(last_success.brands)} brands"

        latest_text = "--"
        latest_meta = "No runs yet"
        if records:
            latest_record = records[-1]
            latest = latest_record.finished_at
            try:
                dt = datetime.fromisoformat(latest)
                latest_text = dt.strftime("%B %d, %Y")
                latest_meta = f"{dt.strftime('%I:%M %p')} · {latest_record.status.capitalize()}"
            except Exception:
                latest_text = latest[:10]
                latest_meta = latest_record.status.capitalize()

        self.stat_db_update.set_values(db_value, db_meta)
        self.stat_last_success.set_values(success_value, success_meta)
        self.stat_recent.set_values(latest_text, latest_meta)


class ProjectPage(QWidget):
    back_requested = Signal()
    refresh_brands_requested = Signal()
    start_requested = Signal(dict)
    review_requested = Signal()

    def __init__(self) -> None:
        super().__init__()
        self.last_output_dir = ""
        self._build()

    def _build(self) -> None:
        root = QVBoxLayout(self)
        root.setContentsMargins(28, 24, 28, 28)
        root.setSpacing(18)

        header_row = QHBoxLayout()
        title_box = QVBoxLayout()
        title = QLabel("Project intake and extraction")
        title.setObjectName("pageTitle")
        subtitle = QLabel("Add every source drawing now. Extraction creates a source-preserved takeoff ready for pricing.")
        subtitle.setObjectName("pageSubtitle")
        title_box.addWidget(title)
        title_box.addWidget(subtitle)
        
        self.back_btn = QPushButton("Back to overview")
        set_button_variant(self.back_btn, "secondary")
        self.back_btn.clicked.connect(self.back_requested.emit)
        
        header_row.addLayout(title_box)
        header_row.addStretch(1)
        header_row.addWidget(self.back_btn, 0, Qt.AlignTop)

        body = QHBoxLayout()
        body.setSpacing(16)

        # LEFT COLUMN: FORM
        config_card = QFrame()
        config_card.setObjectName("card")
        config_layout = QVBoxLayout(config_card)
        config_layout.setContentsMargins(18, 18, 18, 18)
        config_layout.setSpacing(14)
        
        config_title = QLabel("PROJECT MANIFEST")
        config_title.setObjectName("sectionTitle")
        
        form = QFormLayout()
        form.setHorizontalSpacing(16)
        form.setVerticalSpacing(12)
        
        self.project_name = QLineEdit("Project_1")
        self.project_ref = QLineEdit("Client Project")
        self.drawing_list = QListWidget()
        self.drawing_list.setObjectName("drawingManifest")
        self.drawing_list.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.drawing_list.setMinimumHeight(126)
        self.drawing_list.setAccessibleName("Source drawing manifest")
        self.output = QLineEdit(str(PROJECT_ROOT / "output"))
        self.output.setCursorPosition(0)

        draw_btn = QPushButton("Add drawings")
        set_button_variant(draw_btn, "secondary")
        draw_btn.clicked.connect(self._pick_drawing)
        remove_draw_btn = QPushButton("Remove selected")
        set_button_variant(remove_draw_btn, "secondary")
        remove_draw_btn.clicked.connect(self._remove_drawings)
        out_btn = QPushButton("Browse")
        set_button_variant(out_btn, "secondary")
        out_btn.clicked.connect(self._pick_output)

        self.brand_list = QListWidget()
        self.brand_list.setSelectionMode(QAbstractItemView.NoSelection)
        self.brand_list.setMinimumHeight(120)
        
        self.custom_brands = QLineEdit()
        self.custom_brands.setPlaceholderText("Additional brands, comma-separated...")

        self.year = QSpinBox()
        self.year.setRange(2024, 2050)
        self.half = QComboBox()
        self.half.addItems(["H1", "H2"])
        self.pricing_mode = QComboBox()
        self.pricing_mode.addItems(["total", "itemized"])

        form.addRow("Project Name", self.project_name)
        form.addRow("Reference", self.project_ref)
        drawing_box = QWidget()
        drawing_layout = QVBoxLayout(drawing_box)
        drawing_layout.setContentsMargins(0, 0, 0, 0)
        drawing_layout.setSpacing(6)
        drawing_layout.addWidget(self.drawing_list)
        drawing_actions = QHBoxLayout()
        drawing_actions.addWidget(draw_btn)
        drawing_actions.addWidget(remove_draw_btn)
        drawing_actions.addStretch(1)
        self.drawing_count = QLabel("No drawings selected")
        self.drawing_count.setObjectName("sectionDesc")
        drawing_actions.addWidget(self.drawing_count)
        drawing_layout.addLayout(drawing_actions)
        form.addRow("Source drawings", drawing_box)
        form.addRow("Output Folder", self._line_with_button(self.output, out_btn))
        
        brand_row = QVBoxLayout()
        brand_row.setSpacing(8)
        brand_row.addWidget(self.brand_list)
        refresh_link = QPushButton("Refresh brands from database")
        set_button_variant(refresh_link, "ghost")
        refresh_link.setCursor(Qt.PointingHandCursor)
        refresh_link.clicked.connect(self.refresh_brands_requested.emit)
        brand_row.addWidget(refresh_link, 0, Qt.AlignRight)
        
        form.addRow("Available brands", brand_row)
        form.addRow("Additional brands", self.custom_brands)
        form.addRow("Year", self.year)
        form.addRow("Half", self.half)
        form.addRow("Pricing Mode", self.pricing_mode)

        self.start_btn = QPushButton("Start evidence extraction")
        set_button_variant(self.start_btn, "primary")
        self.start_btn.setMinimumHeight(36)
        self.start_btn.clicked.connect(self._emit_start)

        config_layout.addWidget(config_title)
        catalog_note = QLabel(
            "Catalog settings below constrain matching. Extracted components flow directly to pricing unless you exclude them."
        )
        catalog_note.setObjectName("sectionDesc")
        catalog_note.setWordWrap(True)
        config_layout.addWidget(catalog_note)
        config_layout.addLayout(form)
        config_layout.addStretch(1)
        config_layout.addWidget(self.start_btn)

        # RIGHT COLUMN: EXECUTION / LOGS
        run_card = QFrame()
        run_card.setObjectName("card")
        run_layout = QVBoxLayout(run_card)
        run_layout.setContentsMargins(18, 18, 18, 18)
        run_layout.setSpacing(12)
        
        run_header = QHBoxLayout()
        run_title = QLabel("EXTRACTION RUN")
        run_title.setObjectName("sectionTitle")
        self.status_chip = QLabel("Idle")
        self.status_chip.setObjectName("statusChip")
        run_header.addWidget(run_title)
        run_header.addStretch(1)
        run_header.addWidget(self.status_chip)

        self.progress_label = QLabel("Ready to start")
        self.progress_label.setObjectName("sectionDesc")
        self.progress = QProgressBar()
        self.progress.setRange(0, 1)
        self.progress.setValue(0)
        self.progress.setTextVisible(False)
        
        self.log = QTextEdit()
        self.log.setObjectName("terminal")
        self.log.setReadOnly(True)
        self.log.setPlaceholderText("Run details will appear here after extraction starts.")

        run_controls = QHBoxLayout()
        self.clear_log_btn = QPushButton("Clear details")
        set_button_variant(self.clear_log_btn, "secondary")
        self.open_output_btn = QPushButton("Open project folder")
        set_button_variant(self.open_output_btn, "secondary")
        self.open_output_btn.setEnabled(False)
        self.review_btn = QPushButton("Review or price project")
        set_button_variant(self.review_btn, "primary")
        self.review_btn.setEnabled(False)
        self.cancel_btn = QPushButton("Cancel after current drawing")
        set_button_variant(self.cancel_btn, "secondary")
        self.cancel_btn.setEnabled(False)
        
        self.clear_log_btn.clicked.connect(self.log.clear)
        self.open_output_btn.clicked.connect(self._open_last_output)
        self.review_btn.clicked.connect(self.review_requested.emit)
        
        run_controls.addWidget(self.clear_log_btn)
        run_controls.addWidget(self.cancel_btn)
        run_controls.addStretch(1)
        run_controls.addWidget(self.review_btn)
        run_controls.addWidget(self.open_output_btn)

        run_layout.addLayout(run_header)
        run_layout.addWidget(self.progress_label)
        run_layout.addWidget(self.progress)
        run_layout.addSpacing(8)
        run_layout.addWidget(self.log, 1)
        run_layout.addLayout(run_controls)

        body.addWidget(config_card, 5)
        body.addWidget(run_card, 6)

        root.addLayout(header_row)
        root.addLayout(body, 1)

    def _line_with_button(self, line: QLineEdit, button: QPushButton) -> QWidget:
        row = QWidget()
        lay = QHBoxLayout(row)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(8)
        lay.addWidget(line, 1)
        lay.addWidget(button)
        return row

    def _pick_drawing(self) -> None:
        paths, _ = QFileDialog.getOpenFileNames(self, "Choose drawings", str(PROJECT_ROOT), "Drawing Files (*.dwg *.dxf);;All Files (*)")
        if paths:
            existing = {self.drawing_list.item(index).text() for index in range(self.drawing_list.count())}
            for path in paths:
                if path not in existing:
                    self.drawing_list.addItem(path)
            self._update_drawing_count()

    def _remove_drawings(self) -> None:
        for item in self.drawing_list.selectedItems():
            self.drawing_list.takeItem(self.drawing_list.row(item))
        self._update_drawing_count()

    def _update_drawing_count(self) -> None:
        count = self.drawing_list.count()
        self.drawing_count.setText(
            f"{count} drawing{'s' if count != 1 else ''} selected" if count else "No drawings selected"
        )

    def _pick_output(self) -> None:
        folder = QFileDialog.getExistingDirectory(self, "Choose output folder", self.output.text())
        if folder:
            self.output.setText(folder)

    def _emit_start(self) -> None:
        selected = [
            self.brand_list.item(index).text()
            for index in range(self.brand_list.count())
            if self.brand_list.item(index).checkState() == Qt.Checked
        ]
        custom = [x.strip().upper() for x in self.custom_brands.text().split(",") if x.strip()]
        self.start_requested.emit(
            {
                "project_name": self.project_name.text().strip(),
                "project_ref": self.project_ref.text().strip(),
                "drawings": [self.drawing_list.item(index).text() for index in range(self.drawing_list.count())],
                "output": self.output.text().strip(),
                "brands": sorted(set(selected + custom)),
                "year": int(self.year.value()),
                "half": self.half.currentText(),
                "pricing_mode": self.pricing_mode.currentText(),
            }
        )

    def _open_last_output(self) -> None:
        if self.last_output_dir and Path(self.last_output_dir).exists():
            QDesktopServices.openUrl(QUrl.fromLocalFile(self.last_output_dir))

    def set_running(self, running: bool) -> None:
        self.start_btn.setEnabled(not running)
        self.back_btn.setEnabled(not running)
        self.cancel_btn.setEnabled(running)
        if running:
            self.status_chip.setText("Running")
            self.status_chip.setProperty("state", "running")
        else:
            self.status_chip.setText("Idle")
            self.status_chip.setProperty("state", "")
        self.status_chip.style().unpolish(self.status_chip)
        self.status_chip.style().polish(self.status_chip)

    def set_progress(self, step: int, total: int, text: str) -> None:
        self.progress.setRange(0, total)
        self.progress.setValue(step)
        self.progress_label.setText(f"Step {step} of {total}: {text}")

    def set_brands(self, brands: list[str]) -> None:
        self.brand_list.clear()
        for brand in brands:
            item = QListWidgetItem(brand)
            item.setFlags(item.flags() | Qt.ItemIsUserCheckable)
            item.setCheckState(Qt.Checked if brand.upper() == "SCHNEIDER" else Qt.Unchecked)
            self.brand_list.addItem(item)

    def apply_defaults(self, cfg: AppConfig) -> None:
        self.year.setValue(cfg.default_year)
        self.half.setCurrentText(cfg.default_half)
        self.pricing_mode.setCurrentText(cfg.default_pricing_mode)


class HistoryPage(QWidget):
    delete_requested = Signal(int)

    def __init__(self) -> None:
        super().__init__()
        self.records: list[JobRecord] = []
        lay = QVBoxLayout(self)
        lay.setContentsMargins(40, 40, 40, 40)
        lay.setSpacing(24)

        header_box = QVBoxLayout()
        header_box.setSpacing(8)
        title = QLabel("Execution History")
        title.setObjectName("pageTitle")
        subtitle = QLabel("Audit trail of all previous pipeline runs.")
        subtitle.setObjectName("pageSubtitle")
        header_box.addWidget(title)
        header_box.addWidget(subtitle)

        card = QFrame()
        card.setObjectName("card")
        card_lay = QVBoxLayout(card)
        card_lay.setContentsMargins(0, 0, 0, 0)

        self.table = QTableWidget(0, 6)
        self.table.setHorizontalHeaderLabels(["Timestamp", "Project Name", "Brands Executed", "Status", "Output Directory", "Error Message"])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setShowGrid(False)
        self.table.verticalHeader().setVisible(False)
        self.table.setAlternatingRowColors(True)

        card_lay.addWidget(self.table)

        ctrl_row = QHBoxLayout()
        self.delete_btn = QPushButton("Delete Selected Entry")
        set_button_variant(self.delete_btn, "secondary")
        self.delete_btn.setEnabled(False)
        self.open_btn = QPushButton("Open Selected Output")
        set_button_variant(self.open_btn, "primary")
        ctrl_row.addWidget(self.delete_btn)
        ctrl_row.addStretch(1)
        ctrl_row.addWidget(self.open_btn)

        lay.addLayout(header_box)
        lay.addWidget(card, 1)
        lay.addLayout(ctrl_row)

        self.table.itemSelectionChanged.connect(self._sync_action_buttons)
        self.delete_btn.clicked.connect(self._request_delete_selected)
        self.open_btn.clicked.connect(self._open_selected)

    def set_records(self, records: list[JobRecord]) -> None:
        self.records = records
        self.table.setRowCount(len(records))
        for i, rec in enumerate(records):
            # Format timestamp nicely
            try:
                dt = datetime.fromisoformat(rec.finished_at)
                ts = dt.strftime("%Y-%m-%d %H:%M")
            except Exception:
                ts = rec.finished_at[:16]

            self.table.setItem(i, 0, QTableWidgetItem(ts))
            self.table.setItem(i, 1, QTableWidgetItem(rec.project_name))
            self.table.setItem(i, 2, QTableWidgetItem(", ".join(rec.brands)))
            
            status_item = QTableWidgetItem(rec.status.capitalize())
            if rec.status == "success":
                status_item.setForeground(QColor("#059669")) # Emerald
            else:
                status_item.setForeground(QColor("#DC2626")) # Red
            self.table.setItem(i, 3, status_item)
            
            self.table.setItem(i, 4, QTableWidgetItem(rec.output_dir))
            self.table.setItem(i, 5, QTableWidgetItem(rec.error[:120]))
            
        self.table.resizeColumnToContents(0)
        self.table.resizeColumnToContents(1)
        self.table.resizeColumnToContents(2)
        self.table.resizeColumnToContents(3)
        self._sync_action_buttons()

    def _open_selected(self) -> None:
        row = self.table.currentRow()
        if row < 0 or row >= len(self.records):
            return
        folder = Path(self.records[row].output_dir)
        if folder.exists():
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(folder)))

    def _request_delete_selected(self) -> None:
        row = self.table.currentRow()
        if row < 0 or row >= len(self.records):
            return
        self.delete_requested.emit(row)

    def _sync_action_buttons(self) -> None:
        row = self.table.currentRow()
        has_valid_selection = 0 <= row < len(self.records)
        self.delete_btn.setEnabled(has_valid_selection)


class PricingDbPage(QWidget):
    db_changed = Signal()
    import_excel_requested = Signal()

    def __init__(self, db_path_getter, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self._db_path_getter = db_path_getter
        self._columns: list[DbColumnInfo] = []
        self._pk_col: DbColumnInfo | None = None
        self._pk_index: int = -1
        self._original_rows: dict[str, dict[str, object]] = {}
        self._pending_deletes: set[str] = set()
        self._build()
        self.refresh_tables(preserve_selection=False)

    def _build(self) -> None:
        root = QVBoxLayout(self)
        root.setContentsMargins(40, 40, 40, 40)
        root.setSpacing(24)

        header_box = QVBoxLayout()
        header_box.setSpacing(8)
        title = QLabel("Pricing Database Manager")
        title.setObjectName("pageTitle")
        subtitle = QLabel("Read and edit pricing tables directly from the configured SQLite database.")
        subtitle.setObjectName("pageSubtitle")
        subtitle.setWordWrap(True)
        header_box.addWidget(title)
        header_box.addWidget(subtitle)

        controls_card = QFrame()
        controls_card.setObjectName("card")
        controls_layout = QVBoxLayout(controls_card)
        controls_layout.setContentsMargins(20, 16, 20, 16)
        controls_layout.setSpacing(12)

        row1 = QHBoxLayout()
        row1.setSpacing(10)
        row1.addWidget(QLabel("Table"))
        self.table_picker = QComboBox()
        self.table_picker.setMinimumWidth(220)
        self.table_picker.currentTextChanged.connect(self.reload_table)
        row1.addWidget(self.table_picker)

        row1.addSpacing(16)
        row1.addWidget(QLabel("Filter"))
        self.search = QLineEdit()
        self.search.setPlaceholderText("Type to filter visible rows...")
        self.search.textChanged.connect(self._apply_filter)
        row1.addWidget(self.search, 1)

        self.refresh_tables_btn = QPushButton("Refresh Tables")
        set_button_variant(self.refresh_tables_btn, "secondary")
        self.refresh_tables_btn.clicked.connect(lambda _checked=False: self.refresh_tables())
        row1.addWidget(self.refresh_tables_btn)

        self.reload_btn = QPushButton("Reload Data")
        set_button_variant(self.reload_btn, "secondary")
        self.reload_btn.clicked.connect(lambda _checked=False: self.reload_table())
        row1.addWidget(self.reload_btn)

        row2 = QHBoxLayout()
        row2.setSpacing(10)
        self.add_btn = QPushButton("Add Row")
        set_button_variant(self.add_btn, "secondary")
        self.add_btn.clicked.connect(self.add_row)
        row2.addWidget(self.add_btn)

        self.delete_btn = QPushButton("Delete Selected")
        set_button_variant(self.delete_btn, "secondary")
        self.delete_btn.clicked.connect(self.delete_selected_rows)
        row2.addWidget(self.delete_btn)

        self.import_excel_btn = QPushButton("Import Excel")
        set_button_variant(self.import_excel_btn, "primary")
        self.import_excel_btn.clicked.connect(self.import_excel_requested.emit)
        row2.addWidget(self.import_excel_btn)

        row2.addStretch(1)

        self.save_btn = QPushButton("Save Changes")
        set_button_variant(self.save_btn, "primary")
        self.save_btn.clicked.connect(self.save_changes)
        row2.addWidget(self.save_btn)

        controls_layout.addLayout(row1)
        controls_layout.addLayout(row2)

        table_card = QFrame()
        table_card.setObjectName("card")
        table_layout = QVBoxLayout(table_card)
        table_layout.setContentsMargins(0, 0, 0, 0)
        self.table = QTableWidget(0, 0)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setShowGrid(False)
        self.table.verticalHeader().setVisible(False)
        self.table.verticalHeader().setDefaultSectionSize(34)
        self.table.setAlternatingRowColors(True)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.horizontalHeader().setMinimumSectionSize(120)
        self.table.horizontalHeader().setDefaultSectionSize(180)
        self.table.setItemDelegate(WideCellEditorDelegate(parent=self.table))
        self.table.setWordWrap(False)
        table_layout.addWidget(self.table)

        self.status = QLabel("Select a table to start.")
        self.status.setObjectName("sectionDesc")

        root.addLayout(header_box)
        root.addWidget(controls_card)
        root.addWidget(table_card, 1)
        root.addWidget(self.status)

    def _db_path(self) -> Path:
        return Path(str(self._db_path_getter() or "")).expanduser()

    def _set_status(self, text: str) -> None:
        self.status.setText(text)

    def refresh_tables(self, preserve_selection: bool = True) -> None:
        previous = self.table_picker.currentText().strip() if preserve_selection else ""
        db_path = self._db_path()
        if not db_path.exists():
            self.table_picker.blockSignals(True)
            self.table_picker.clear()
            self.table_picker.blockSignals(False)
            self._clear_editor()
            self._set_status(f"Database file not found: {db_path}")
            return

        try:
            conn = sqlite3.connect(str(db_path))
            names = list_sqlite_tables(conn)
            conn.close()
        except Exception as exc:
            self._clear_editor()
            self._set_status(f"Failed to read tables: {exc}")
            return

        preferred = ["products", "ppl_editions", "manufacturers", "product_keywords", "term_aliases", "match_log"]
        ordered = [name for name in preferred if name in names] + [name for name in names if name not in preferred]

        self.table_picker.blockSignals(True)
        self.table_picker.clear()
        self.table_picker.addItems(ordered)
        self.table_picker.blockSignals(False)

        if not ordered:
            self._clear_editor()
            self._set_status("No tables found in database.")
            return

        if previous and previous in ordered:
            self.table_picker.setCurrentText(previous)
        else:
            self.table_picker.setCurrentIndex(0)
        self.reload_table()

    def _clear_editor(self) -> None:
        self._columns = []
        self._pk_col = None
        self._pk_index = -1
        self._original_rows.clear()
        self._pending_deletes.clear()
        self.table.clear()
        self.table.setRowCount(0)
        self.table.setColumnCount(0)

    def reload_table(self, _selected: str | None = None) -> None:
        table_name = self.table_picker.currentText().strip()
        if not table_name:
            self._clear_editor()
            self._set_status("Select a table to start.")
            return

        db_path = self._db_path()
        if not db_path.exists():
            self._clear_editor()
            self._set_status(f"Database file not found: {db_path}")
            return

        try:
            conn = sqlite3.connect(str(db_path))
            conn.row_factory = sqlite3.Row
            self._columns = get_table_columns(conn, table_name)
            self._pk_col = next((c for c in sorted(self._columns, key=lambda c: c.pk_pos) if c.pk_pos > 0), None)
            quoted_table = quote_ident(table_name)
            total_count = int(conn.execute(f"SELECT COUNT(*) FROM {quoted_table}").fetchone()[0])
            order_clause = f" ORDER BY {quote_ident(self._pk_col.name)}" if self._pk_col else ""
            rows = conn.execute(f"SELECT * FROM {quoted_table}{order_clause}").fetchall()
            conn.close()
        except Exception as exc:
            self._clear_editor()
            self._set_status(f"Failed to load table '{table_name}': {exc}")
            return

        self._pending_deletes.clear()
        self._original_rows.clear()
        self._pk_index = -1
        headers = [c.name for c in self._columns]

        self.table.clear()
        self.table.setColumnCount(len(headers))
        self.table.setHorizontalHeaderLabels(headers)
        self.table.setRowCount(len(rows))

        if self._pk_col is not None:
            self._pk_index = headers.index(self._pk_col.name)

        for row_idx, row in enumerate(rows):
            for col_idx, col in enumerate(self._columns):
                raw = row[col.name]
                item = QTableWidgetItem("" if raw is None else str(raw))
                if self._pk_col and col.name == self._pk_col.name and raw is not None:
                    item.setFlags(item.flags() & ~Qt.ItemIsEditable)
                self.table.setItem(row_idx, col_idx, item)
            if self._pk_col is not None:
                raw_pk = row[self._pk_col.name]
                if raw_pk is not None:
                    self._original_rows[str(raw_pk)] = {
                        c.name: row[c.name]
                        for c in self._columns
                        if c.name != self._pk_col.name
                    }

        self._apply_filter()
        self._tune_table_columns()
        msg = f"Loaded {len(rows)} row(s) from '{table_name}'."
        if total_count != len(rows):
            msg += f" (Table total: {total_count})"
        self._set_status(msg)

    def _tune_table_columns(self) -> None:
        if self.table.columnCount() == 0:
            return
        self.table.resizeColumnsToContents()
        for col in range(self.table.columnCount()):
            width = self.table.columnWidth(col)
            tuned = min(max(width + 24, 160), 520)
            self.table.setColumnWidth(col, tuned)

    def _apply_filter(self, _text: str | None = None) -> None:
        needle = self.search.text().strip().lower()
        rows = self.table.rowCount()
        cols = self.table.columnCount()
        if not needle:
            for row in range(rows):
                self.table.setRowHidden(row, False)
            return
        for row in range(rows):
            visible = False
            for col in range(cols):
                item = self.table.item(row, col)
                if item and needle in item.text().lower():
                    visible = True
                    break
            self.table.setRowHidden(row, not visible)

    def add_row(self, _checked: bool = False) -> None:
        if not self._columns:
            self._set_status("No table is loaded.")
            return
        row = self.table.rowCount()
        self.table.insertRow(row)
        for col_idx, col in enumerate(self._columns):
            item = QTableWidgetItem("")
            if self._pk_col and col.name == self._pk_col.name:
                item.setFlags(item.flags() & ~Qt.ItemIsEditable)
            self.table.setItem(row, col_idx, item)
        self.table.setCurrentCell(row, 0)
        self._apply_filter()
        self._set_status("New row added. Fill values and click Save Changes.")

    def delete_selected_rows(self, _checked: bool = False) -> None:
        selected_rows = sorted({idx.row() for idx in self.table.selectedIndexes()}, reverse=True)
        if not selected_rows:
            self._set_status("No rows selected.")
            return

        for row in selected_rows:
            if self._pk_col is not None and self._pk_index >= 0:
                pk_item = self.table.item(row, self._pk_index)
                pk_text = pk_item.text().strip() if pk_item else ""
                if pk_text:
                    self._pending_deletes.add(pk_text)
            self.table.removeRow(row)

        self._set_status(
            f"Marked {len(selected_rows)} row(s) for deletion. Click Save Changes to commit."
        )

    def _parse_cell(self, text: str, col: DbColumnInfo, row_num: int) -> object:
        raw = text.strip()
        if raw == "":
            return None

        dtype = (col.declared_type or "").upper()
        try:
            if "INT" in dtype:
                return int(raw)
            if any(token in dtype for token in ("REAL", "FLOA", "DOUB")):
                return float(raw)
        except ValueError as exc:
            raise ValueError(
                f"Row {row_num}: '{col.name}' expects {col.declared_type or 'a numeric value'}, got '{text}'."
            ) from exc
        return raw

    def _coerce_pk_value(self, raw_pk: str) -> object:
        if self._pk_col is None:
            return raw_pk
        dtype = (self._pk_col.declared_type or "").upper()
        if "INT" in dtype:
            return int(raw_pk)
        return raw_pk

    def save_changes(self, _checked: bool = False) -> None:
        table_name = self.table_picker.currentText().strip()
        if not table_name or not self._columns:
            self._set_status("No table is loaded.")
            return
        if self._pk_col is None or self._pk_index < 0:
            QMessageBox.warning(
                self,
                "Save Not Supported",
                "This table has no primary key. Editing is disabled for safety.",
            )
            return

        db_path = self._db_path()
        if not db_path.exists():
            QMessageBox.warning(self, "Database Missing", f"Database file not found:\n{db_path}")
            return

        quoted_table = quote_ident(table_name)
        quoted_pk = quote_ident(self._pk_col.name)
        inserted = 0
        updated = 0
        deleted = 0

        try:
            conn = sqlite3.connect(str(db_path))
            conn.execute("BEGIN")

            for row in range(self.table.rowCount()):
                pk_item = self.table.item(row, self._pk_index)
                pk_text = pk_item.text().strip() if pk_item else ""

                row_values: dict[str, object] = {}
                for col_idx, col in enumerate(self._columns):
                    if col.name == self._pk_col.name:
                        continue
                    item = self.table.item(row, col_idx)
                    value = self._parse_cell(item.text() if item else "", col, row + 1)
                    if value is None and col.not_null:
                        raise ValueError(f"Row {row + 1}: '{col.name}' is required.")
                    row_values[col.name] = value

                if not pk_text:
                    if all(v is None for v in row_values.values()):
                        continue
                    insert_cols = [name for name, value in row_values.items() if value is not None]
                    if not insert_cols:
                        continue
                    placeholders = ", ".join("?" for _ in insert_cols)
                    sql = (
                        f"INSERT INTO {quoted_table} ("
                        + ", ".join(quote_ident(name) for name in insert_cols)
                        + f") VALUES ({placeholders})"
                    )
                    conn.execute(sql, [row_values[name] for name in insert_cols])
                    inserted += 1
                    continue

                if pk_text not in self._original_rows:
                    raise ValueError(
                        f"Row {row + 1}: primary key was changed. Editing primary key values is not supported."
                    )

                previous = self._original_rows[pk_text]
                changed_cols = [name for name in row_values if row_values[name] != previous.get(name)]
                if changed_cols:
                    set_clause = ", ".join(f"{quote_ident(name)} = ?" for name in changed_cols)
                    params = [row_values[name] for name in changed_cols]
                    params.append(self._coerce_pk_value(pk_text))
                    conn.execute(
                        f"UPDATE {quoted_table} SET {set_clause} WHERE {quoted_pk} = ?",
                        params,
                    )
                    updated += 1

            for raw_pk in sorted(self._pending_deletes):
                result = conn.execute(
                    f"DELETE FROM {quoted_table} WHERE {quoted_pk} = ?",
                    (self._coerce_pk_value(raw_pk),),
                )
                if result.rowcount and result.rowcount > 0:
                    deleted += int(result.rowcount)

            conn.commit()
            conn.close()
        except Exception as exc:
            try:
                conn.rollback()  # type: ignore[name-defined]
                conn.close()  # type: ignore[name-defined]
            except Exception:
                pass
            QMessageBox.warning(self, "Save Failed", str(exc))
            return

        self.reload_table()
        self.db_changed.emit()
        self._set_status(
            f"Saved successfully. Inserted: {inserted}, Updated: {updated}, Deleted: {deleted}."
        )


class MainWindow(QMainWindow):
    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle("Technicat Automation")
        self.resize(1360, 860)

        self.cfg = load_config()
        default_db = PROJECT_ROOT / "pricing.db"
        if not Path(self.cfg.db_path).exists() and default_db.exists():
            self.cfg.db_path = str(default_db)
            save_config(self.cfg)
        self.history = load_history()
        self.worker_thread: QThread | None = None
        self.worker: ProjectWorker | None = None
        self.import_thread: QThread | None = None
        self.import_worker: ExcelImportWorker | None = None
        self.import_progress: QProgressDialog | None = None

        self._build_ui()
        self.project_page.apply_defaults(self.cfg)
        self.refresh_brands()
        self.history_page.set_records(list(reversed(self.history)))
        self._refresh_home_overview()
        
        self.nav_home.setChecked(True)

    def _build_ui(self) -> None:
        central = QWidget()
        main_layout = QHBoxLayout(central)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)

        # Build Sidebar
        sidebar = QFrame()
        sidebar.setObjectName("sidebar")
        sidebar.setFixedWidth(216)
        sidebar_lay = QVBoxLayout(sidebar)
        sidebar_lay.setContentsMargins(14, 18, 14, 18)
        sidebar_lay.setSpacing(8)

        logo_lay = QVBoxLayout()
        logo_lay.setSpacing(8)

        logo_mark = QLabel()
        logo_mark.setObjectName("appLogoMark")
        logo_mark.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        logo_mark.setFixedSize(164, 58)
        logo_pixmap = self._load_company_logo(logo_mark.size())
        if logo_pixmap is not None:
            logo_mark.setPixmap(logo_pixmap)
        else:
            logo_mark.setText("TG")
            logo_mark.setObjectName("appLogo")

        logo = QLabel("Technicat Automation")
        logo.setObjectName("appLogo")
        logo_sub = QLabel("")
        logo_sub.setObjectName("appLogoSubtitle")
        logo_sub.setVisible(False)
        logo_lay.addWidget(logo_mark)
        logo_lay.addWidget(logo)
        logo_lay.addWidget(logo_sub)

        sidebar_lay.addLayout(logo_lay)
        sidebar_lay.addSpacing(24)

        self.nav_group = QButtonGroup(self)
        self.nav_group.setExclusive(True)

        self.nav_home = self._create_nav_button("Overview", self.show_home)
        self.nav_project = self._create_nav_button("Project intake", self.show_project)
        self.nav_history = self._create_nav_button("Run history", self.show_history)
        self.nav_pricing_db = self._create_nav_button("Pricing database", self.show_pricing_db)

        sidebar_lay.addWidget(self.nav_home)
        sidebar_lay.addWidget(self.nav_project)
        sidebar_lay.addWidget(self.nav_history)
        sidebar_lay.addWidget(self.nav_pricing_db)
        sidebar_lay.addStretch(1)

        settings_btn = QPushButton("Preferences")
        settings_btn.setObjectName("navBtn")
        settings_btn.setCursor(Qt.PointingHandCursor)
        settings_btn.clicked.connect(self.open_settings)
        sidebar_lay.addWidget(settings_btn)

        # Build Stack
        self.stack = QStackedWidget()
        self.home_page = HomePage()
        self.project_page = ProjectPage()
        self.history_page = HistoryPage()
        self.pricing_db_page = PricingDbPage(lambda: self.cfg.db_path)
        self.stack.addWidget(self.home_page)
        self.stack.addWidget(self.project_page)
        self.stack.addWidget(self.history_page)
        self.stack.addWidget(self.pricing_db_page)

        # Assembly
        main_layout.addWidget(sidebar)
        main_layout.addWidget(self.stack, 1)
        self.setCentralWidget(central)

        # Wire up cross-page events
        self.home_page.new_project_requested.connect(self._nav_to_project)
        self.home_page.history_requested.connect(self._nav_to_history)
        self.home_page.pricing_db_requested.connect(self._nav_to_pricing_db)
        self.home_page.import_excel_requested.connect(self.import_excel_to_db)
        self.project_page.back_requested.connect(self._nav_to_home)
        self.project_page.refresh_brands_requested.connect(self.refresh_brands)
        self.project_page.start_requested.connect(self.start_project)
        self.project_page.cancel_btn.clicked.connect(self.cancel_project)
        self.project_page.review_requested.connect(self.open_project_review)
        self.pricing_db_page.db_changed.connect(self.refresh_brands)
        self.pricing_db_page.db_changed.connect(self._refresh_home_overview)
        self.pricing_db_page.import_excel_requested.connect(self.import_excel_to_db)
        self.history_page.delete_requested.connect(self.delete_history_entry)

    def _load_company_logo(self, target_size: QSize) -> QPixmap | None:
        candidates = [
            PROJECT_ROOT / "technicat_logo.png",
            PROJECT_ROOT / "technicat_logo.jpg",
            PROJECT_ROOT / "technicat_logo.jpeg",
            PROJECT_ROOT / "assets" / "technicat_logo.png",
            PROJECT_ROOT / "assets" / "technicat_logo.jpg",
            PROJECT_ROOT / "assets" / "technicat_logo.jpeg",
            SOURCE_ROOT / "technicat_logo.png",
            SOURCE_ROOT / "technicat_logo.jpg",
            SOURCE_ROOT / "technicat_logo.jpeg",
            SOURCE_ROOT / "assets" / "technicat_logo.png",
            SOURCE_ROOT / "assets" / "technicat_logo.jpg",
            SOURCE_ROOT / "assets" / "technicat_logo.jpeg",
            BUNDLE_ROOT / "technicat_logo.png",
            BUNDLE_ROOT / "technicat_logo.jpg",
            BUNDLE_ROOT / "technicat_logo.jpeg",
            BUNDLE_ROOT / "assets" / "technicat_logo.png",
            BUNDLE_ROOT / "assets" / "technicat_logo.jpg",
            BUNDLE_ROOT / "assets" / "technicat_logo.jpeg",
        ]
        for path in candidates:
            if not path.exists():
                continue
            pixmap = QPixmap(str(path))
            if pixmap.isNull():
                continue
            return pixmap.scaled(target_size, Qt.KeepAspectRatio, Qt.SmoothTransformation)
        return None

    def _create_nav_button(self, text: str, callback) -> QPushButton:
        btn = QPushButton(text)
        btn.setObjectName("navBtn")
        btn.setCheckable(True)
        btn.setCursor(Qt.PointingHandCursor)
        btn.clicked.connect(callback)
        self.nav_group.addButton(btn)
        return btn

    def _refresh_home_overview(self) -> None:
        self.home_page.set_overview(self.history, self.cfg.db_path)

    def show_home(self) -> None:
        self.stack.setCurrentWidget(self.home_page)

    def show_project(self) -> None:
        self.stack.setCurrentWidget(self.project_page)

    def show_history(self) -> None:
        self.history_page.set_records(list(reversed(self.history)))
        self.stack.setCurrentWidget(self.history_page)

    def show_pricing_db(self) -> None:
        self.pricing_db_page.refresh_tables()
        self.stack.setCurrentWidget(self.pricing_db_page)

    def open_project_review(self) -> None:
        if not self.project_page.last_output_dir:
            return
        project_path = Path(self.project_page.last_output_dir) / CANONICAL_FILENAME
        if not project_path.exists():
            QMessageBox.warning(self, "Review unavailable", f"Project file not found: {project_path}")
            return
        ProjectReviewDialog(project_path, self.cfg.db_path, self).exec()

    def _nav_to_project(self) -> None:
        self.nav_project.setChecked(True)
        self.show_project()

    def _nav_to_history(self) -> None:
        self.nav_history.setChecked(True)
        self.show_history()

    def _nav_to_pricing_db(self) -> None:
        self.nav_pricing_db.setChecked(True)
        self.show_pricing_db()

    def _nav_to_home(self) -> None:
        self.nav_home.setChecked(True)
        self.show_home()

    def delete_history_entry(self, reversed_row_index: int) -> None:
        displayed = list(reversed(self.history))
        if reversed_row_index < 0 or reversed_row_index >= len(displayed):
            return

        rec = displayed[reversed_row_index]
        answer = QMessageBox.question(
            self,
            "Delete History Entry",
            (
                "Delete the selected history entry?\n\n"
                f"Project: {rec.project_name}\n"
                f"Timestamp: {rec.finished_at}\n\n"
                "This action cannot be undone."
            ),
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No,
        )
        if answer != QMessageBox.Yes:
            return

        source_index = len(self.history) - 1 - reversed_row_index
        if source_index < 0 or source_index >= len(self.history):
            return

        self.history.pop(source_index)
        save_history(self.history)
        self._refresh_home_overview()
        self.history_page.set_records(list(reversed(self.history)))

    def open_settings(self) -> None:
        dlg = SettingsDialog(self.cfg, self)
        if dlg.exec() != QDialog.Accepted:
            return
        self.cfg = dlg.to_config()
        save_config(self.cfg)
        self.project_page.apply_defaults(self.cfg)
        self.refresh_brands()
        self.pricing_db_page.refresh_tables()
        self._refresh_home_overview()

    def import_excel_to_db(self) -> None:
        if self.import_thread is not None:
            QMessageBox.information(self, "Import In Progress", "An Excel import is already running.")
            return

        dlg = ImportExcelDialog(self.cfg, self)
        if dlg.exec() != QDialog.Accepted:
            return

        payload = dlg.payload()
        excel_path = Path(payload["excel_file"]).expanduser()
        db_path = Path(payload["db_path"]).expanduser()
        script_path = RUNTIME_CODES_DIR / "import.py"

        try:
            if not script_path.exists():
                raise ValueError(f"Missing script: {script_path}")
            if not excel_path.exists():
                raise ValueError(f"Excel file not found: {excel_path}")
            if not IS_FROZEN and not Path(self.cfg.python_executable).exists():
                raise ValueError(f"Python executable not found: {self.cfg.python_executable}")

            db_path.parent.mkdir(parents=True, exist_ok=True)
            self._sync_import_key_env()

            env = os.environ.copy()
            if (self.cfg.gemini_api_key or "").strip():
                env["GEMINI_API_KEY"] = self.cfg.gemini_api_key.strip()
            if (self.cfg.gemini_model or "").strip():
                env["GEMINI_API_MODEL"] = self.cfg.gemini_model.strip()
            env["PYTHONUTF8"] = "1"
            env["PYTHONIOENCODING"] = "utf-8"

            script_args = [
                "--file", str(excel_path),
                "--manufacturer", str(payload["manufacturer"]),
                "--year", str(payload["year"]),
                "--half", str(payload["half"]),
                "--currency", str(payload["currency"]),
                "--db", str(db_path),
            ]
            if payload["force"]:
                script_args.append("--force")
        except Exception as exc:
            QMessageBox.warning(self, "Import Failed", str(exc))
            return

        self.import_progress = QProgressDialog("Importing Excel into database...", "", 0, 0, self)
        self.import_progress.setWindowTitle("Import Running")
        self.import_progress.setCancelButton(None)
        self.import_progress.setMinimumDuration(0)
        self.import_progress.setWindowModality(Qt.WindowModal)
        self.import_progress.show()

        self.import_thread = QThread()
        self.import_worker = ExcelImportWorker(
            python_executable=self.cfg.python_executable,
            script_path=script_path,
            script_args=script_args,
            env=env,
            db_path=db_path,
            manufacturer=str(payload["manufacturer"]).upper(),
            year=int(payload["year"]),
            half=str(payload["half"]),
        )
        self.import_worker.moveToThread(self.import_thread)
        self.import_thread.started.connect(self.import_worker.run)
        self.import_worker.finished.connect(self.on_import_excel_finished)
        self.import_worker.failed.connect(self.on_import_excel_failed)
        self.import_worker.finished.connect(self.import_thread.quit)
        self.import_worker.failed.connect(self.import_thread.quit)
        self.import_thread.finished.connect(self._cleanup_import_worker)
        self.import_thread.start()

    def _sync_import_key_env(self) -> None:
        legacy_env = PROJECT_ROOT / ".venv" / "key.env"
        existing: dict[str, str] = {}
        if legacy_env.exists():
            for line in legacy_env.read_text(encoding="utf-8", errors="replace").splitlines():
                raw = line.strip()
                if not raw or raw.startswith("#") or "=" not in raw:
                    continue
                key, value = raw.split("=", 1)
                existing[key.strip()] = value.strip()

        key = (self.cfg.gemini_api_key or "").strip() or existing.get("GEMINI_API_KEY", "")
        model = (self.cfg.gemini_model or "").strip() or existing.get("GEMINI_API_MODEL", DEFAULT_GEMINI_MODEL)
        if not key:
            return

        existing["GEMINI_API_KEY"] = key
        existing["GEMINI_API_MODEL"] = model
        legacy_env.parent.mkdir(parents=True, exist_ok=True)
        body = "\n".join(f"{k}={v}" for k, v in existing.items()).strip()
        legacy_env.write_text(body + "\n", encoding="utf-8")

    def on_import_excel_finished(self, payload: dict) -> None:
        if self.import_progress is not None:
            self.import_progress.close()

        db_path = payload.get("db_path", "")
        summary = payload.get("summary", {})
        self.cfg.db_path = str(db_path)
        save_config(self.cfg)
        self.refresh_brands()
        self._refresh_home_overview()
        self.pricing_db_page.refresh_tables(preserve_selection=False)
        self.nav_pricing_db.setChecked(True)
        self.show_pricing_db()

        QMessageBox.information(
            self,
            "Import Complete",
            (
                "Excel was imported successfully.\n\n"
                f"DB switched to: {db_path}\n"
                f"Edition products imported: {summary.get('edition_products', 0)}\n"
                f"Total products in DB: {summary.get('total_products', 0)}"
            ),
        )

    def on_import_excel_failed(self, error: str) -> None:
        if self.import_progress is not None:
            self.import_progress.close()
        QMessageBox.warning(self, "Import Failed", error)

    def _cleanup_import_worker(self) -> None:
        self.import_worker = None
        self.import_thread = None
        if self.import_progress is not None:
            self.import_progress.close()
        self.import_progress = None

    def refresh_brands(self) -> None:
        brands = fetch_brands(self.cfg.db_path)
        self.project_page.set_brands(brands or DEFAULT_BRANDS)

    def start_project(self, payload: dict) -> None:
        try:
            self._validate(payload)
            self.project_page.log.clear()
            self.project_page.set_running(True)
            self.project_page.set_progress(0, 1, "Initializing...")
            
            # Print a clean preamble to the terminal
            preamble = (
                f"=== WORKFLOW INITIALIZED ===\n"
                f"Project: {payload['project_name']}\n"
                f"Model: {self.cfg.gemini_model}\n"
                f"Target Database: {Path(self.cfg.db_path).name}\n"
                f"===========================\n"
            )
            self.project_page.log.append(preamble)

            self.worker_thread = QThread()
            self.worker = ProjectWorker(
                cfg=self.cfg,
                project_name=payload["project_name"],
                project_ref=payload["project_ref"] or payload["project_name"],
                drawing_paths=[Path(value) for value in payload["drawings"]],
                output_dir=Path(payload["output"]),
                brands=payload["brands"],
                year=payload["year"],
                half=payload["half"],
                pricing_mode=payload["pricing_mode"],
            )
            self.worker.moveToThread(self.worker_thread)
            self.worker_thread.started.connect(self.worker.run)
            self.worker.log.connect(self.project_page.log.append)
            self.worker.step.connect(self.project_page.set_progress)
            self.worker.finished.connect(self.on_finished)
            self.worker.failed.connect(self.on_failed)
            self.worker.cancelled.connect(self.on_cancelled)
            self.worker.finished.connect(self.worker_thread.quit)
            self.worker.failed.connect(self.worker_thread.quit)
            self.worker.cancelled.connect(self.worker_thread.quit)
            self.worker_thread.finished.connect(self._cleanup_worker)
            self.worker_thread.start()
        except Exception as exc:
            QMessageBox.warning(self, "Validation Error", str(exc))

    def _validate(self, payload: dict) -> None:
        if not payload["project_name"].strip():
            raise ValueError("Project name is required.")
        if not payload["drawings"] or any(not Path(value).exists() for value in payload["drawings"]):
            raise ValueError("Select one or more valid DWG/DXF files.")
        if not Path(payload["output"]).exists():
            raise ValueError("Output folder does not exist.")
        if not payload["brands"]:
            raise ValueError("Select at least one brand.")
        if not Path(self.cfg.db_path).exists():
            raise ValueError(f"Pricing DB not found: {self.cfg.db_path}")
        legacy_env = PROJECT_ROOT / ".venv" / "key.env"
        has_cfg_key = bool((self.cfg.gemini_api_key or "").strip())
        if not has_cfg_key and not legacy_env.exists():
            raise ValueError("Gemini API key is required. Set it in Preferences or .venv/key.env.")
        self._validate_pipeline_runtime()

    def _validate_pipeline_runtime(self) -> None:
        if IS_FROZEN:
            try:
                __import__("openpyxl")
                __import__("ezdxf")
                __import__("matplotlib")
                __import__("google.genai")
            except Exception as exc:
                raise ValueError(
                    "Bundled runtime is missing pipeline dependencies "
                    f"(openpyxl, ezdxf, matplotlib, google-genai). Details: {exc}"
                ) from exc
            return

        cmd = [
            self.cfg.python_executable,
            "-c",
            "import openpyxl, ezdxf, matplotlib, google.genai; print('ok')",
        ]
        try:
            proc = subprocess.run(
                cmd,
                cwd=str(PROJECT_ROOT),
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=20,
            )
        except Exception as exc:
            raise ValueError(f"Failed to validate pipeline runtime: {exc}") from exc
        if proc.returncode != 0:
            details = (proc.stdout or "").strip().splitlines()
            tail = details[-1] if details else "unknown dependency error"
            raise ValueError(
                "Configured Python is missing pipeline dependencies "
                f"(openpyxl, ezdxf, matplotlib, google-genai). Details: {tail}"
            )

    def on_finished(self, ctx: dict) -> None:
        self.project_page.set_running(False)
        self.project_page.set_progress(1, 1, "Extraction ready for review")
        output_dir = str(Path(ctx["raw_output"]).parent)
        self.project_page.last_output_dir = output_dir
        self.project_page.open_output_btn.setEnabled(True)
        self.project_page.review_btn.setEnabled(True)

        self.history.append(
            JobRecord(
                finished_at=ctx["finished_at"],
                project_name=ctx["project_name"],
                brands=ctx["brands"],
                status="success",
                output_dir=ctx["output_dir"],
            )
        )
        save_history(self.history)
        self._refresh_home_overview()
        
        self.project_page.log.append("\n=== EXTRACTION COMPLETE — REVIEW REQUIRED BEFORE PRICING ===")

    def cancel_project(self) -> None:
        if self.worker_thread is None:
            return
        self.worker_thread.requestInterruption()
        self.project_page.cancel_btn.setEnabled(False)
        self.project_page.progress_label.setText(
            "Cancellation requested — finishing the current drawing safely"
        )
        self.project_page.log.append(
            "Cancellation requested. The current drawing will finish before the run stops."
        )

    def on_cancelled(self, ctx: dict) -> None:
        self.project_page.set_running(False)
        self.project_page.set_progress(0, 1, "Run cancelled safely")
        self.project_page.log.append(
            "\n=== EXTRACTION CANCELLED — COMPLETED DRAWINGS WERE SAVED ==="
        )
        self.history.append(JobRecord(
            finished_at=ctx.get("finished_at", datetime.now().replace(microsecond=0).isoformat()),
            project_name=ctx.get("project_name", ""),
            brands=ctx.get("brands", []),
            status="cancelled",
            output_dir=ctx.get("output_dir", ""),
        ))
        save_history(self.history)
        self._refresh_home_overview()

    def on_failed(self, error: str, ctx: dict) -> None:
        self.project_page.set_running(False)
        self.project_page.set_progress(1, 1, "Execution Failed")
        self.project_page.log.append(f"\n[ERROR] Workflow failed: {error}")

        self.history.append(
            JobRecord(
                finished_at=ctx.get("finished_at", datetime.now().replace(microsecond=0).isoformat()),
                project_name=ctx.get("project_name", ""),
                brands=ctx.get("brands", []),
                status="failed",
                output_dir=ctx.get("output_dir", ""),
                error=error,
            )
        )
        save_history(self.history)
        self._refresh_home_overview()
        QMessageBox.critical(self, "Execution Failed", f"The pipeline encountered an error:\n\n{error}")

    def _cleanup_worker(self) -> None:
        self.worker = None
        self.worker_thread = None


def main() -> int:
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    app.setStyleSheet(APP_STYLESHEET)
    window = MainWindow()
    window.show()
    return app.exec()

if __name__ == "__main__":
    raise SystemExit(main())
