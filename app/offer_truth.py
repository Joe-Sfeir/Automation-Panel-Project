from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any


@dataclass
class OfferTruthRow:
    panel_reference: str
    description: str
    manufacturer: str
    quantity: int
    scope: str
    category: str
    poles: str = "-"
    amps: str = "-"
    curve: str = "-"
    breaking_capacity: str = "-"
    sensitivity: str = "-"
    spd_type: str = "-"
    model: str = "-"
    ct_ratio: str = "-"


PANEL_ACCESSORY = re.compile(
    r"\b(?:panel|side panel|plastron|rail|door|enclosure|cabinet|cubicle)\b", re.I
)
TERMINAL = re.compile(r"\bterminal\b", re.I)


def _poles(text: str) -> str:
    normalized = text.upper().replace("×", "X")
    match = re.search(r"\b([1234])\s*P(?:\s*\+\s*N)?\b", normalized)
    if match:
        return f"{match.group(1)}P+N" if re.search(rf"\b{match.group(1)}\s*P\s*\+\s*N", normalized) else f"{match.group(1)}P"
    match = re.search(r"\b([1234])\s*X\s*\d+(?:\.\d+)?\s*A\b", normalized)
    return f"{match.group(1)}P" if match else "-"


def _amps(text: str, category: str) -> str:
    if category == "CT":
        return "-"
    normalized = text.upper().replace("×", "X")
    matches = re.findall(r"(?<![\d/])(\d+(?:\.\d+)?)\s*A\b", normalized)
    return f"{matches[-1]}A" if matches else "-"


def _category(text: str) -> str:
    value = text.upper()
    if value.startswith("RCCB"):
        return "RCCB"
    if value.startswith("RCBO"):
        return "RCBO"
    if value.startswith("MCCB"):
        return "MCCB"
    if value.startswith("MCB"):
        return "MCB"
    if "SURGE" in value or re.search(r"\bSPD\b", value):
        return "SPD"
    if value.startswith("CT "):
        return "CT"
    if "POWER METER" in value:
        return "METER"
    if "PILOT LAMP" in value:
        return "PILOT_LAMP"
    if "FUSE HOLDER" in value:
        return "FUSE_HOLDER"
    if "CONTACTOR" in value or value.startswith("ATS "):
        return "CONTACTOR"
    if "SWITCH" in value or value.startswith("MTS "):
        return "SWITCH"
    if "CONTROL RELAY" in value or "VOLTAGE RELAY" in value:
        return "RELAY"
    if "SELECTOR" in value:
        return "SELECTOR_SWITCH"
    if "TIMER" in value:
        return "TIMER"
    if value == "PLC":
        return "PLC"
    if "POWER SUPPLY" in value:
        return "POWER_SUPPLY"
    return "OTHER"


def parse_offer_description(panel_reference: str, description: str, manufacturer: str, quantity: int) -> OfferTruthRow:
    text = " ".join(str(description).split())
    if TERMINAL.search(text):
        scope = "calculated_terminal"
        category = "TERMINAL"
    elif PANEL_ACCESSORY.search(text):
        scope = "commercial_accessory"
        category = "ENCLOSURE_ACCESSORY"
    else:
        scope = "sld_visible"
        category = _category(text)
    normalized = text.upper()
    curve = "-"
    curve_match = re.search(r"(?:,|\bCURVE)\s*([BCD])\s*(?:,|$)", normalized)
    if curve_match:
        curve = curve_match.group(1)
    capacity_match = re.search(r"\b(\d+(?:\.\d+)?)\s*KA\b", normalized)
    sensitivity_match = re.search(r"\b(\d+(?:\.\d+)?)\s*MA\b", normalized)
    ct_match = re.search(r"\bCT\s*(\d+\s*/\s*\d+)\b", normalized)
    spd_match = re.search(r"\bTYPE\s*([123])\b", normalized)
    model = "-"
    for pattern in (
        r"\b(INS\d+)\b", r"\b(I?C\d+[A-Z]?)\b", r"\b(NSX\d+[A-Z]?)\b",
        r"\b(LC[12]-[A-Z0-9]+)\b",
    ):
        match = re.search(pattern, normalized)
        if match:
            model = match.group(1)
            break
    return OfferTruthRow(
        panel_reference=panel_reference,
        description=text,
        manufacturer=str(manufacturer or "").strip().upper(),
        quantity=max(0, int(quantity)),
        scope=scope,
        category=category,
        poles=_poles(text),
        amps=_amps(text, category),
        curve=curve,
        breaking_capacity=f"{capacity_match.group(1)}kA" if capacity_match and category != "SPD" else "-",
        sensitivity=f"{sensitivity_match.group(1)}mA" if sensitivity_match else "-",
        spd_type=f"TYPE {spd_match.group(1)}" if spd_match else "-",
        model=model,
        ct_ratio=ct_match.group(1).replace(" ", "") if ct_match else "-",
    )


def import_offer_truth(workbook_path: str | Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to import benchmark offer workbooks") from exc
    source = Path(workbook_path)
    workbook = load_workbook(source, data_only=True, read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    headers: list[int] = []
    for index, row in enumerate(rows[:-1]):
        panel = row[2] if len(row) > 2 else None
        total = row[5] if len(row) > 5 else None
        next_label = rows[index + 1][4] if len(rows[index + 1]) > 4 else None
        if panel and isinstance(total, (int, float)) and str(next_label or "").strip().upper() == "TTC:":
            headers.append(index)
    panels = []
    for position, start in enumerate(headers):
        end = headers[position + 1] if position + 1 < len(headers) else len(rows)
        reference = " ".join(str(rows[start][2]).split())
        items = []
        for row in rows[start + 2:end]:
            description = row[2] if len(row) > 2 else None
            manufacturer = row[3] if len(row) > 3 else None
            quantity = row[5] if len(row) > 5 else None
            if description and isinstance(quantity, (int, float)):
                items.append(asdict(parse_offer_description(reference, description, manufacturer, int(quantity))))
        panels.append({"reference": reference, "items": items})
    workbook.close()
    return {"source_workbook": str(source), "sheet": sheet.title, "panels": panels}
