import json

from openpyxl import load_workbook

from app.exporting import export_pricing_json


def test_excel_export_respects_visible_columns(tmp_path):
    source = tmp_path / "pricing.json"
    source.write_text(json.dumps({
        "project_reference": "REF-1", "grand_ht": 20, "tva_rate": 0.1, "tva": 2, "grand_ttc": 22,
        "panels": [{"panel_reference": "P1", "items": [{
            "description": "MCB", "chosen_reference": "ABC", "quantity": 2,
            "total_price": 20, "review_status": "confirmed",
        }]}],
    }), encoding="utf-8")
    output = export_pricing_json(source, tmp_path / "offer.xlsx", ["panel", "description", "quantity", "total_price"])
    sheet = load_workbook(output, data_only=True).active
    assert [sheet.cell(6, column).value for column in range(1, 5)] == ["Panel", "Description", "Qty", "Total"]
    assert [sheet.cell(7, column).value for column in range(1, 5)] == ["P1", "MCB", 2, 20]
